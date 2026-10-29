"""Pantallas extra del panel: tablero, cola de asesores, evento en vivo, QR, mapa, reporte ejecutivo, scoring, mapeo."""

import json
from datetime import date, datetime, time, timedelta
from io import BytesIO

from django.conf import settings
from django.contrib import admin, messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.exceptions import PermissionDenied
from django.db.models import Count, F, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone

from . import catalogos
from .models import (etiqueta_segmento, ACTIVOS, Preparatoria, Configuracion, Cuestionario, Evento, Formulario, PosibleCoincidencia, Pregunta, Prospecto,
                     Segmento, Seguimiento, TipoSeguimiento)
from .servicios import auditoria, importacion, inteligencia, reportes


def _ctx(request, titulo, **kw):
    return admin.site.each_context(request) | {'title': titulo, **kw}


def _permiso(request, *permisos):
    if not any(request.user.has_perm(f'prospectos.{p}') for p in permisos):
        raise PermissionDenied


def _fecha(valor, defecto):
    try:
        return date.fromisoformat(valor) if valor else defecto
    except ValueError:
        return defecto


def _evento(request, pk):
    e = get_object_or_404(Evento.objects.select_related('preparatoria', 'promotor__perfil'), pk=pk)
    _permiso(request, 'view_evento')
    if not request.user.is_superuser and request.user.has_perm('prospectos.ver_eventos_propios') and e.promotor_id != request.user.pk:
        raise PermissionDenied
    return e


# --------------------------------------------------------------------------- tablero (portada del panel)
COLORES_SEGMENTO = {
    'listo_para_convertir': '#30b14c', 'sensible_beca': '#f08a00', 'horario_flexible': '#2ba6d6',
    'indeciso_carrera': '#a550c9', 'no_contactar': '#ff3b30',
}


def _corta(texto, n=24):
    texto = str(texto)
    return texto if len(texto) <= n else texto[:n - 1].rstrip() + '…'


def _grafica(filas, etiqueta='etiqueta', valor='total', color='var(--color-primary-600)', **extra):
    return json.dumps({'labels': [_corta(f[etiqueta]) for f in filas],
                       'datasets': [{'data': [f[valor] for f in filas], 'backgroundColor': color, 'borderColor': color,
                                     'borderRadius': 4, 'maxBarThickness': 16, **extra}]})


FUENTE = {'family': 'system-ui, -apple-system, "Segoe UI", Roboto, sans-serif', 'size': 12}


def _opciones_horizontales():
    return json.dumps({'indexAxis': 'y', 'maintainAspectRatio': False, 'plugins': {'legend': {'display': False}},
                       'scales': {'x': {'beginAtZero': True, 'ticks': {'precision': 0, 'font': FUENTE}, 'border': {'display': False}},
                                  'y': {'grid': {'display': False}, 'ticks': {'font': FUENTE}, 'border': {'display': False}}}})


def _opciones_linea():
    return json.dumps({'maintainAspectRatio': False, 'plugins': {'legend': {'display': False}},
                       'scales': {'x': {'grid': {'display': False}, 'ticks': {'font': FUENTE}, 'border': {'display': False}},
                                  'y': {'beginAtZero': True, 'ticks': {'precision': 0, 'font': FUENTE}, 'border': {'display': False}}}})


def _opciones_dona():
    return json.dumps({'maintainAspectRatio': False, 'cutout': '64%',
                       'plugins': {'legend': {'position': 'right', 'labels': {'boxWidth': 10, 'usePointStyle': True, 'font': FUENTE}}}})


PERIODOS = [(7, '7 días'), (30, '30 días'), (90, '90 días'), (365, '12 meses')]
COLORES_ANILLO = ['var(--c-verde)', 'var(--c-naranja)', 'var(--c-cian)', 'var(--c-morado)']


def _base_tablero(request):
    qs = Prospecto.objects.reportables()
    u = request.user
    if u.has_perm('prospectos.ver_propios') and not u.is_superuser:
        qs = qs.filter(promotor=u)  # el promotor solo ve lo suyo
    return qs


def _actividad(qs, dias: int, hoy):
    """Registros por día (o por semana si el periodo es largo) y crecimiento contra el periodo anterior."""
    from collections import Counter
    inicio = hoy - timedelta(days=dias - 1)
    desde = timezone.make_aware(datetime.combine(inicio, time.min))
    # Se agrupa en Python (hora local): MariaDB sin tablas de zonas horarias devuelve NULL con CONVERT_TZ.
    por_dia = Counter(timezone.localtime(f).date() for f in qs.filter(creado_en__gte=desde).values_list('creado_en', flat=True))
    dias_lista = [inicio + timedelta(days=i) for i in range(dias)]
    nombres = ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']
    if dias <= 31:
        etiquetas = [nombres[d.weekday()] if dias <= 7 else d.strftime('%d/%m') for d in dias_lista]
        valores = [por_dia.get(d, 0) for d in dias_lista]
    else:
        etiquetas, valores = [], []
        for i in range(0, dias, 7):
            bloque = dias_lista[i:i + 7]
            etiquetas.append(bloque[0].strftime('%d/%m'))
            valores.append(sum(por_dia.get(d, 0) for d in bloque))
    actual = sum(por_dia.values())
    previo_desde = timezone.make_aware(datetime.combine(inicio - timedelta(days=dias), time.min))
    previo = qs.filter(creado_en__gte=previo_desde, creado_en__lt=desde).count()
    crecimiento = round(100 * (actual - previo) / previo) if previo else None
    return {'etiquetas': etiquetas, 'valores': valores, 'total': actual, 'previo': previo, 'crecimiento': crecimiento}


def contexto_tablero(request, context=None) -> dict:
    """Portada del panel (DASHBOARD_CALLBACK de unfold): primero a quién llamar hoy; luego el resumen del periodo."""
    ctx = context if context is not None else {}
    u = request.user
    hoy = timezone.localdate()
    try:
        dias = int(request.GET.get('dias', 30))
    except ValueError:
        dias = 30
    dias = dias if dias in dict(PERIODOS) else 30
    desde = timezone.make_aware(datetime.combine(hoy - timedelta(days=dias - 1), time.min))
    ctx.update({
        'hoy': hoy, 'dias': dias, 'periodos': PERIODOS, 'periodo_txt': dict(PERIODOS)[dias].replace('12 meses', 'año'),
        'puede_reportes': u.has_perm('prospectos.ver_reportes'),
        'eventos_hoy': Evento.objects.filter(fecha=hoy).select_related('preparatoria'),
    })
    tarjetas = []
    if u.has_perm('prospectos.change_prospecto'):
        from .admin import iniciales
        base = _cola_base()
        llamar = list(base[:8])
        for p in llamar:
            p.iniciales = iniciales(p.nombre_aspirante)
        ctx['llamar'] = llamar
        ctx['llamar_total'] = base.count()
        listos = Prospecto.objects.utilizables().filter(estatus__in=ACTIVOS, segmentos__icontains=Segmento.LISTO.value).count()
        tareas = badge_tareas(request) or 0
        coinc = PosibleCoincidencia.objects.filter(estado='pendiente').count()
        tarjetas += [
            {'valor': listos, 'etiqueta': 'Listos para inscribirse', 'icono': 'rocket_launch', 'tono': 'verde',
             'url': reverse('prospectos:cola') + '?segmento=listo_para_convertir',
             'ayuda': 'Pidieron información por WhatsApp o correo y tienen prioridad alta. Contáctalos primero.'},
            {'valor': tareas, 'etiqueta': 'Llamadas pendientes', 'icono': 'task_alt', 'tono': 'ambar',
             'url': reverse('prospectos:cola') + '?tareas=1',
             'ayuda': 'Tareas abiertas de los asesores (incluye las automáticas por días sin contacto).'},
            {'valor': coinc, 'etiqueta': 'Posibles duplicados', 'icono': 'join_inner', 'tono': 'magenta',
             'url': reverse('admin:prospectos_posiblecoincidencia_changelist'),
             'ayuda': 'Dos registros que parecen la misma persona. Alguien debe decidir si se unen.'},
        ]
        ctx['formularios_error'] = Formulario.objects.filter(activo=True).exclude(ultimo_error='')
    if u.has_perm('prospectos.view_prospecto'):
        esperando = badge_esperando(request) or 0
        if esperando:
            tarjetas.append({'valor': esperando, 'etiqueta': 'Esperando permiso', 'icono': 'hourglass_top', 'tono': 'ambar',
                             'url': reverse('prospectos:respuestas') + '?estado=cuarentena',
                             'ayuda': 'Respuestas de menores sin el permiso del tutor: no se usan hasta que llegue.'})
    if u.is_superuser or u.has_perm('prospectos.change_formulario'):
        pasos = primeros_pasos()
        if not all(p['hecho'] for p in pasos) and not Configuracion.obtener('pasos.ocultos', False):
            ctx['primeros_pasos'] = pasos
            ctx['pasos_hechos'] = sum(p['hecho'] for p in pasos)
            ctx['pasos_listos'] = ', '.join(p['titulo'][0].lower() + p['titulo'][1:] for p in pasos if p['hecho'])
    if u.is_superuser:
        from .models import Respaldo
        ultimo = Respaldo.objects.order_by('-creado_en').first()
        if ultimo and ultimo.estado != 'ok':
            ctx['respaldo_error'] = ultimo
    if ctx['puede_reportes']:
        qs = _base_tablero(request)
        periodo = qs.filter(creado_en__gte=desde)
        total_periodo = periodo.count()
        tarjetas.insert(0, {'valor': total_periodo, 'etiqueta': 'Aspirantes nuevos', 'resumen': True, 'icono': 'groups',
                            'tono': 'azul', 'url': reverse('admin:prospectos_prospecto_changelist'),
                            'ayuda': 'Personas distintas registradas en el periodo (alumno y tutor cuentan como uno).'})
        actividad = _actividad(qs, dias, hoy)

        # Anillos: % del periodo en cada segmento clave.
        from collections import Counter
        conteo = Counter()
        for segs in periodo.values_list('segmentos', flat=True):
            conteo.update(segs or [])
        claves = [Segmento.LISTO, Segmento.BECA, Segmento.HORARIO, Segmento.INDECISO]
        anillos = [{'etiqueta': Segmento(c).label, 'pct': round(100 * conteo.get(c.value, 0) / total_periodo) if total_periodo else 0,
                    'color': COLORES_ANILLO[i]} for i, c in enumerate(claves)]
        piden_info = periodo.filter(canal_informacion__in=['whatsapp', 'correo', 'si']).count()

        # Barras por edad.
        edades = Counter()
        for e in periodo.exclude(edad=None).values_list('edad', flat=True):
            edades['≤15' if e <= 15 else '19+' if e >= 19 else str(e)] += 1
        orden = ['≤15', '16', '17', '18', '19+']
        tope = max([edades.get(k, 0) for k in orden] or [1]) or 1
        barras_edad = [{'etiqueta': f'{k} años', 'total': edades.get(k, 0), 'pct': round(100 * edades.get(k, 0) / tope)} for k in orden]

        # Mapa: preparatorias con coordenadas y aspirantes del periodo (por plantel).
        puntos = reportes.Metricas.cobertura(periodo)
        ctx['sin_ubicar'] = reportes.Metricas.sin_ubicar().filter(total__gt=0).count()

        m = reportes.Metricas(hoy - timedelta(days=dias - 1), hoy)
        if u.has_perm('prospectos.ver_propios') and not u.is_superuser:
            m.base = lambda: periodo
        embudo = m.embudo()
        base_embudo = max(1, embudo[0]['total'] if embudo else 1)
        for e in embudo:
            e['pct'] = round(100 * e['total'] / base_embudo, 1)
        ranking = reportes.Metricas.ranking_promotores(
            timezone.make_aware(datetime.combine(hoy, time.min)), timezone.make_aware(datetime.combine(hoy, time.max)))
        maximo = max([r['total'] for r in ranking] or [1])
        for r in ranking:
            r['pct'] = round(100 * r['total'] / maximo)
        carreras = m.carreras(6)
        tope_c = max([c['total'] for c in carreras] or [1])
        for c in carreras:
            c['pct'] = round(100 * c['total'] / tope_c)

        recientes = []
        if u.has_perm('prospectos.view_prospecto'):
            from .admin import iniciales
            for p in qs.select_related('preparatoria', 'carrera').order_by('-creado_en')[:6]:
                recientes.append({'id': p.pk, 'nombre': p.nombre_aspirante, 'iniciales': iniciales(p.nombre_aspirante),
                                  'detalle': p.carrera.nombre if p.carrera else (p.preparatoria.nombre if p.preparatoria else 'Sin datos'),
                                  'cuando': p.creado_en, 'score': p.score})
        ctx.update({
            'recientes': recientes, 'embudo': embudo, 'ranking_hoy': ranking[:5], 'carreras': carreras,
            'info_pct': round(100 * piden_info / total_periodo) if total_periodo else 0,
            'actividad': actividad, 'barras_edad': barras_edad, 'anillos': anillos,
            'datos_tablero': {'actividad': actividad, 'anillos': anillos, 'puntos': puntos, 'tiles': TILES()},
        })
    ctx['tarjetas'] = tarjetas
    return ctx


def badge_tareas(request):
    if not request.user.has_perm('prospectos.change_prospecto'):
        return None
    n = Seguimiento.objects.filter(tipo=TipoSeguimiento.TAREA, completada_en__isnull=True,
                                   prospecto__consentimiento=True, prospecto__anonimizado_en__isnull=True).count()
    return n or None


def badge_coincidencias(request):
    if not request.user.has_perm('prospectos.resolver_coincidencias'):
        return None
    return PosibleCoincidencia.objects.filter(estado='pendiente').exclude(motivo__startswith='Familia').count() or None


def badge_esperando(request):
    """Respuestas retenidas por falta de permiso del tutor (las que más confunden: 'contesté y no aparece')."""
    if not request.user.has_perm('prospectos.view_prospecto'):
        return None
    from .models import EstadoRespuesta
    return _respuestas_visibles(request).filter(estado=EstadoRespuesta.CUARENTENA).count() or None


def primeros_pasos() -> list[dict]:
    """Lista de lo que falta para usar el programa con datos reales. Cada paso dice cómo resolverlo."""
    from django.contrib.auth import get_user_model
    from .models import Respuesta
    from .servicios import fuentes
    forms = {t: Formulario.objects.filter(activo=True, tipo=t).exclude(google_form_id__startswith='demo-').first()
             for t in (Cuestionario.ALUMNO, Cuestionario.TUTOR)}
    alumno = forms[Cuestionario.ALUMNO]
    codigos = set((alumno.mapeo_campos or {}).values()) if alumno else set()
    lista_forms = reverse('admin:prospectos_formulario_changelist')
    return [
        {'titulo': 'Conectar con Google', 'hecho': fuentes.credenciales_listas(), 'url': None,
         'como': f'Pon la llave de la cuenta de servicio en la carpeta de datos ({settings.DATA_DIR}) como '
                 f'google-service-account.json (o ejecuta «{settings.COMANDO} conectar_google»).'},
        {'titulo': 'Registrar el Form del alumno', 'hecho': bool(alumno), 'url': lista_forms,
         'como': 'Formularios de Google → Agregar → pega la URL de edición (la que termina en /edit).'},
        {'titulo': 'Pregunta de edad en el Form', 'hecho': 'edad' in codigos,
         'url': reverse('prospectos:mapeo', args=[alumno.pk]) if alumno else lista_forms,
         'como': 'Sin la edad no se sabe quién es menor: todo se queda «Esperando permiso». Pon validación de número en el Form.'},
        {'titulo': 'Aviso de privacidad', 'hecho': 'consentimiento_aviso_privacidad' in codigos or bool(inteligencia.aviso_en_formulario()),
         'url': reverse('prospectos:configuracion') + '#permisos',
         'como': 'Agrega al Form «He leído y acepto el aviso de privacidad» o, si el aviso ya aparece en el Form, márcalo en Reglas de permiso.'},
        {'titulo': 'Recibir la primera respuesta', 'hecho': Respuesta.objects.exclude(formulario__google_form_id__startswith='demo-').exists(),
         'url': reverse('prospectos:respuestas'),
         'como': 'Contesta tu propio Form. Se importa solo cada minuto mientras la ventana negra esté abierta.'},
        {'titulo': 'Ubicar las preparatorias en el mapa',
         'hecho': Preparatoria.objects.exists() and not reportes.Metricas.sin_ubicar().filter(total__gt=0).exists(),
         'url': reverse('prospectos:ubicar'), 'como': 'Las que escriben los alumnos se dan de alta solas; ubícalas con un clic.'},
        {'titulo': 'Crear usuarios del equipo', 'hecho': get_user_model().objects.filter(is_active=True).count() > 1,
         'url': reverse('admin:auth_user_changelist'), 'como': 'Una cuenta por persona, con su rol. No compartan contraseñas.'},
        {'titulo': 'Guardar copia de config.json fuera de esta computadora', 'hecho': bool(Configuracion.obtener('pasos.config_copiado', False)),
         'url': None, 'marcable': 'config_copiado',
         'como': 'Tiene las llaves de cifrado. Si se pierde, los teléfonos y correos guardados no se recuperan nunca.'},
    ]


@staff_member_required
def pasos(request):
    """Marca a mano los pasos que el programa no puede comprobar solo, u oculta la lista."""
    if request.method == 'POST' and (request.user.is_superuser or request.user.has_perm('prospectos.change_formulario')):
        accion = request.POST.get('accion')
        if accion == 'config_copiado':
            Configuracion.guardar('pasos.config_copiado', True, request.user)
        elif accion == 'ocultar':
            Configuracion.guardar('pasos.ocultos', True, request.user)
        elif accion == 'mostrar':
            Configuracion.guardar('pasos.ocultos', False, request.user)
    return redirect(request.POST.get('siguiente') or 'admin:index')


@staff_member_required
def ayuda(request):
    """«¿Cómo funciona?»: el programa explicado sin tecnicismos."""
    p = settings.PROSPECTOS
    pesos = inteligencia.pesos()
    total = sum(float(v) for v in pesos.values()) or 1
    reglas = [(NOMBRES_PESOS.get(k, k), round(float(v) / total * 100)) for k, v in pesos.items() if float(v) > 0]
    puede_config = request.user.is_superuser or request.user.has_perm('prospectos.change_formulario')
    return render(request, 'prospectos/ayuda.html', _ctx(
        request, '¿Cómo funciona?', reglas=reglas,
        listo=inteligencia.umbral('listo_para_convertir_score'), indeciso=inteligencia.umbral('indeciso_desde_carreras'),
        cuarentena=p.get('CUARENTENA_DIAS', 30), sin_contacto=p.get('DIAS_SIN_CONTACTO', 5), retencion=p.get('RETENCION_MESES', 18),
        pasos=primeros_pasos() if puede_config else None, pasos_ocultos=Configuracion.obtener('pasos.ocultos', False)))


def huella() -> str:
    """Identifica la carpeta de datos de esta instancia (para no confundir la de pruebas con la real)."""
    import hashlib
    return hashlib.sha256(str(settings.DATA_DIR.resolve()).encode()).hexdigest()[:12]


def logo(request):
    """Logo opcional de la institución (datos/logo.png). Público: se usa también en el login."""
    from django.http import FileResponse, Http404
    if not settings.LOGO_ARCHIVO.exists():
        raise Http404
    return FileResponse(open(settings.LOGO_ARCHIVO, 'rb'), content_type='image/png')


# --------------------------------------------------------------------------- cola de asesores
def _cola_base():
    """Aspirantes que se pueden contactar, del más prioritario al que lleva más tiempo sin contacto."""
    return (Prospecto.objects.utilizables().filter(estatus__in=ACTIVOS, no_contactar=False)
            .select_related('preparatoria', 'carrera', 'asesor')
            .annotate(tareas=Count('seguimientos', filter=Q(seguimientos__tipo=TipoSeguimiento.TAREA,
                                                            seguimientos__completada_en__isnull=True)))
            .order_by('-score', F('ultimo_contacto_en').asc(nulls_first=True)))


@staff_member_required
def cola(request):
    _permiso(request, 'change_prospecto')
    qs = _cola_base()
    f = request.GET
    if f.get('mios'):
        qs = qs.filter(asesor=request.user)
    if f.get('tareas'):
        qs = qs.filter(tareas__gt=0)
    if f.get('segmento'):
        qs = qs.filter(segmentos__icontains=f'"{f["segmento"]}"')
    if f.get('brecha'):
        qs = qs.filter(tiene_brecha=True)
    if f.get('plantel') in catalogos.PLANTELES:
        qs = qs.filter(plantel=f['plantel'])
    if f.get('canal') in ('whatsapp', 'correo'):
        qs = qs.filter(canal_informacion=f['canal'])
    if f.get('q'):
        qs = qs.filter(nombre_aspirante__icontains=f['q'])
    filas = list(qs[:200])
    from .admin import iniciales
    for p in filas:
        todos = p.segmentos or []
        p.segmentos_det = [(s, etiqueta_segmento(s)) for s in todos[:3]]
        p.segmentos_mas = len(todos) - len(p.segmentos_det)
        p.iniciales = iniciales(p.nombre_aspirante)
    no_quieren = Prospecto.objects.utilizables().filter(estatus__in=ACTIVOS, no_contactar=True).count()
    return render(request, 'prospectos/cola.html', _ctx(request, 'A quién llamar hoy', filas=filas, total=qs.count(),
                                                          segmentos=[s for s in Segmento.choices if s[0] != Segmento.NO_CONTACTAR],
                                                          planteles=catalogos.PLANTELES_CORTO.items(), f=f, no_quieren=no_quieren,
                                                          hay_tutor=reportes.formulario_activo(Cuestionario.TUTOR)))


# --------------------------------------------------------------------------- evento en vivo y QR
@staff_member_required
def evento_vivo(request, pk):
    e = _evento(request, pk)
    return render(request, 'prospectos/evento_vivo.html', _ctx(request, f'En vivo · {e}', evento=e,
                                                                 qr_alumno=reportes.qr_general(e, Cuestionario.ALUMNO),
                                                                 qr_tutor=reportes.qr_general(e, Cuestionario.TUTOR)))


@staff_member_required
def evento_vivo_datos(request, pk):
    e = _evento(request, pk)
    c = reportes.Metricas.conteo_evento(e)
    hoy = timezone.localdate()
    ranking = reportes.Metricas.ranking_promotores(timezone.make_aware(datetime.combine(hoy, time.min)),
                                                   timezone.make_aware(datetime.combine(hoy, time.max)))
    return JsonResponse({'alumnos': c['alumnos'], 'tutores': c['tutores'], 'meta': e.meta_registros,
                         'ultima': timezone.localtime(c['ultima']).strftime('%H:%M:%S') if c['ultima'] else None,
                         'ranking': ranking[:10]})


FIRMA_OMISION = {'nombre': '', 'cargo': 'Coordinación de Admisiones', 'telefono': '', 'correo': '',
                 'lugar': 'Cuautitlán Izcalli, Estado de México'}


def _firma():
    return {**FIRMA_OMISION, **(Configuracion.obtener('eventos.firma', {}) or {}), 'logo': settings.LOGO_ARCHIVO.exists()}


@staff_member_required
def evento_solicitud(request, pk):
    """Carta para pedir a la preparatoria el permiso de la visita (imprimir o guardar como PDF)."""
    e = _evento(request, pk)
    if request.method == 'POST':
        _permiso(request, 'change_evento')
        if request.POST.get('accion') == 'firma':
            Configuracion.guardar('eventos.firma', {k: (request.POST.get(k) or '').strip()[:150] for k in FIRMA_OMISION}, request.user)
            auditoria.registrar('firma_solicitudes', None, {}, request=request)
            messages.success(request, 'Datos de quien firma guardados para todas las solicitudes.')
        elif request.POST.get('accion') == 'solicitado' and e.estado == 'por_solicitar':
            e.estado = 'solicitado'
            e.save(update_fields=['estado'])
            messages.success(request, 'Evento marcado como solicitado.')
        return redirect(request.path)
    return render(request, 'prospectos/evento_solicitud.html', {
        'e': e, 'firma': _firma(), 'hoy': timezone.localdate(), 'puede_editar': request.user.has_perm('prospectos.change_evento'),
        'promotor': e.promotor, 'title': f'Solicitud · {e.preparatoria}'})


@staff_member_required
def evento_reporte(request, pk):
    """Reporte final del evento, para imprimir o guardar como PDF."""
    from .servicios import eventos
    e = _evento(request, pk)
    con_nombres = request.GET.get('nombres') == '1' and request.user.has_perm('prospectos.view_prospecto')
    if request.method == 'POST' and request.POST.get('accion') == 'reatribuir':
        _permiso(request, 'change_evento')
        n = eventos.reatribuir(desde=e.fecha)
        messages.success(request, f'Se asignaron {n} respuestas a sus eventos por preparatoria y fecha.')
        return redirect(request.get_full_path())
    d = eventos.reporte(e, con_nombres)
    if con_nombres:
        auditoria.registrar('reporte_evento_con_nombres', None, {'evento': e.codigo}, request=request)
    return render(request, 'prospectos/evento_reporte.html', {
        'e': e, 'd': d, 'firma': _firma(), 'hoy': timezone.localdate(), 'con_nombres': con_nombres,
        'puede_nombres': request.user.has_perm('prospectos.view_prospecto'), 'title': f'Reporte · {e.preparatoria}'})


@staff_member_required
def prepa_qr(request, pk):
    """QR permanente de una preparatoria (carteles y todas sus visitas)."""
    _permiso(request, 'view_preparatoria')
    prepa = get_object_or_404(Preparatoria, pk=pk)
    qrs = [('Cuestionario del alumno', reportes.qr_permanente(prepa, Cuestionario.ALUMNO)),
           ('Encuesta para mamá, papá o tutor', reportes.qr_permanente(prepa, Cuestionario.TUTOR))]
    return render(request, 'prospectos/prepa_qr.html', _ctx(request, f'QR permanente · {prepa}', prepa=prepa, qrs=qrs))


@staff_member_required
def evento_qr(request, pk):
    e = _evento(request, pk)
    qrs = [('Cuestionario del ALUMNO', 'school', reportes.qr_general(e, Cuestionario.ALUMNO)),
           ('Cuestionario del TUTOR', 'family_restroom', reportes.qr_general(e, Cuestionario.TUTOR))]
    return render(request, 'prospectos/evento_qr.html', _ctx(request, f'QR · {e.preparatoria}', evento=e, qrs=qrs))


@staff_member_required
def tarjetas(request, pk):
    e = _evento(request, pk)
    try:
        cantidad = int(request.GET.get('cantidad', 0))
    except ValueError:
        cantidad = 0
    lista = reportes.tarjetas(e, cantidad) if cantidad else []
    if lista:
        auditoria.registrar('generar_tarjetas', e, {'cantidad': len(lista)}, request=request)
    faltan = [t for t in (Cuestionario.ALUMNO, Cuestionario.TUTOR) if not reportes.formulario_activo(t)]
    return render(request, 'prospectos/tarjetas.html', _ctx(request, f'Tarjetas familiares · {e}', evento=e,
                                                              tarjetas=lista, cantidad=cantidad or 20, faltan=faltan))


# --------------------------------------------------------------------------- mapa
@staff_member_required
def mapa(request):
    _permiso(request, 'ver_reportes')
    plantel = request.GET.get('plantel') if request.GET.get('plantel') in catalogos.PLANTELES else None
    base = Prospecto.objects.reportables()
    if plantel:
        base = base.filter(plantel=plantel)
    puntos = [p for p in reportes.Metricas.cobertura(base) if p['total'] or not plantel]
    for p in puntos:
        p['tasa'] = round(100 * p['total'] / p['estimados'], 1) if p['estimados'] else None
    puntos.sort(key=lambda p: -p['total'])
    sin_ubicar = list(reportes.Metricas.sin_ubicar().filter(total__gt=0)[:50])
    return render(request, 'prospectos/mapa.html', _ctx(
        request, 'Mapa de cobertura', puntos=puntos, plantel=plantel, planteles=catalogos.PLANTELES_CORTO.items(),
        sin_ubicar=sin_ubicar, puede_ubicar=request.user.has_perm('prospectos.change_preparatoria'), tiles=TILES()))


OSM = 'https://tile.openstreetmap.org/{z}/{x}/{y}.png'


def TILES() -> dict:  # noqa: N802 - constante calculada
    """
    Fondo del mapa. Por defecto OpenStreetMap (gratis, sin llave; oscurecido con CSS para el tema).
    CARTO dejó de servir sin llave («API KEY REQUIRED»): si alguien lo tiene configurado, se cambia por OSM.
    `MAPA_TILES` en config.json permite otro proveedor (URL con {z}/{x}/{y}); `MAPA_TILES_ATRIBUCION` su crédito.
    """
    url = (settings.CONFIG.get('MAPA_TILES') or settings.CONFIG.get('MAPA_TILES_OSCURO') or '').strip()
    if not url or 'basemaps.cartocdn.com' in url:
        url = OSM
    es_osm = 'tile.openstreetmap.org' in url
    g = mapa_google()
    return {'url': url, 'oscurecer': es_osm,
            'atribucion': '&copy; <a href="https://www.openstreetmap.org/copyright" target="_blank">OpenStreetMap</a>'
                          if es_osm else settings.CONFIG.get('MAPA_TILES_ATRIBUCION', ''),
            'google': {'llave': g['llave'], 'tipo': g['tipo']} if g['llave'] else None}


def mapa_google() -> dict:
    """Llave de Google Maps (Configuración → Mapa). Es una llave de navegador: se protege restringiéndola, no ocultándola."""
    g = Configuracion.obtener('mapa.google') or {}
    return {'llave': (g.get('llave') or settings.CONFIG.get('GOOGLE_MAPS_API_KEY') or '').strip(),
            'tipo': g.get('tipo') if g.get('tipo') in ('roadmap', 'hybrid') else 'roadmap'}


@staff_member_required
def configurar_mapa(request):
    """Elegir el mapa: Google Maps (con llave) u OpenStreetMap (sin llave)."""
    if not request.user.is_superuser:
        raise PermissionDenied
    if request.method == 'POST':
        import re
        llave = (request.POST.get('llave') or '').strip()
        if llave and not re.fullmatch(r'AIza[0-9A-Za-z_-]{35}', llave):
            messages.error(request, 'Eso no parece una llave de Google Maps (empiezan con «AIza» y tienen 39 caracteres).')
            return redirect(request.path)
        tipo = request.POST.get('tipo') if request.POST.get('tipo') in ('roadmap', 'hybrid') else 'roadmap'
        Configuracion.guardar('mapa.google', {'llave': llave, 'tipo': tipo}, request.user)
        auditoria.registrar('configurar_mapa', None, {'google': bool(llave), 'tipo': tipo}, request=request)
        messages.success(request, 'Guardado. Revisa el mapa de prueba de abajo.' if llave else 'Se usará OpenStreetMap (sin llave).')
        return redirect(request.path)
    g = mapa_google()
    return render(request, 'prospectos/configurar_mapa.html', _ctx(
        request, 'Mapa: Google Maps u OpenStreetMap', g=g, tiles=TILES(),
        origen=request.build_absolute_uri('/').rstrip('/'), desde_config=bool(settings.CONFIG.get('GOOGLE_MAPS_API_KEY'))))


# --------------------------------------------------------------------------- familias por unir
def badge_familias(request):
    if not request.user.has_perm('prospectos.resolver_coincidencias'):
        return None
    from .servicios import familias
    return familias.pendientes().count() or None


@staff_member_required
def familias_por_unir(request):
    """Encuestas de padres ↔ cuestionarios de alumnos: confirmar con un clic, o unir a mano buscando por nombre."""
    _permiso(request, 'resolver_coincidencias')
    from .servicios import familias
    from .servicios.inteligencia import NoFusionable, fusionar, rechazar
    if request.method == 'POST':
        accion = request.POST.get('accion')
        try:
            if accion in ('unir', 'no'):
                c = get_object_or_404(PosibleCoincidencia, pk=request.POST.get('propuesta'), estado='pendiente')
                if accion == 'unir':
                    a, f = (c.candidato, c.prospecto) if c.prospecto.tiene_tutor and not c.prospecto.tiene_alumno else (c.prospecto, c.candidato)
                    fusionar(a, f, request.user)
                    messages.success(request, f'Listo: {a.nombre_aspirante} ya tiene a su familia en la ficha.')
                else:
                    rechazar(c, request.user)
            elif accion == 'manual':
                f = get_object_or_404(Prospecto, pk=request.POST.get('familia'), tiene_tutor=True, tiene_alumno=False)
                a = get_object_or_404(Prospecto, pk=request.POST.get('alumno'), tiene_alumno=True, tiene_tutor=False)
                fusionar(a, f, request.user)
                messages.success(request, f'Listo: {a.nombre_aspirante} ya tiene a su familia en la ficha.')
            elif accion == 'buscar_todo':
                r = familias.buscar_para_todos()
                messages.success(request, f"Se unieron {r['unidas']} familias sin duda y quedaron {r['propuestas']} propuestas nuevas para revisar.")
        except NoFusionable as e:
            messages.error(request, str(e))
        return redirect(request.get_full_path())

    propuestas = list(familias.pendientes().order_by('-similitud')[:100])
    for c in propuestas:  # mostrar siempre: izquierda la familia, derecha el alumno
        c.familia, c.alumno = (c.prospecto, c.candidato) if c.prospecto.tiene_tutor and not c.prospecto.tiene_alumno else (c.candidato, c.prospecto)
        c.puntos = round(c.similitud * 100)
        c.razones = c.motivo.split(': ', 1)[-1]
    con_propuesta = {x for c in propuestas for x in (c.prospecto_id, c.candidato_id)}
    vivos = Prospecto.objects.filter(fusionado_en__isnull=True, anonimizado_en__isnull=True)
    solas = vivos.filter(tiene_tutor=True, tiene_alumno=False).exclude(pk__in=con_propuesta).order_by('-creado_en')[:60]
    elegida = request.GET.get('familia')
    q = (request.GET.get('q') or '').strip()
    resultados = []
    if elegida and q:
        resultados = list(vivos.filter(tiene_alumno=True, tiene_tutor=False, nombre_aspirante__icontains=q)
                          .select_related('preparatoria').order_by('nombre_aspirante')[:20])
    return render(request, 'prospectos/familias.html', _ctx(
        request, 'Familias por unir', propuestas=propuestas, solas=solas, elegida=int(elegida) if (elegida or '').isdigit() else None,
        q=q, resultados=resultados,
        total_solas=vivos.filter(tiene_tutor=True, tiene_alumno=False).exclude(pk__in=con_propuesta).count(),
        total_alumnos_solos=vivos.filter(tiene_alumno=True, tiene_tutor=False).count()))


# --------------------------------------------------------------------------- uso desde otras computadoras
@staff_member_required
def otras_computadoras(request):
    """Activar el acceso por red (HTTPS) para que Windows y Mac entren con el navegador."""
    if not request.user.is_superuser:
        raise PermissionDenied
    from .servicios import red
    if request.method == 'POST':
        activar = request.POST.get('accion') == 'activar'
        red.guardar_activa(settings.CONFIG_FILE, activar)
        if activar:
            red.asegurar_certificados(settings.DATA_DIR)
        auditoria.registrar('acceso_red', None, {'activa': activar}, request=request)
        messages.success(request, 'Guardado. Cierra la ventana del programa y vuélvelo a abrir para que tome el cambio.')
        return redirect(request.path)
    import json as _json
    guardado = red.activa(_json.loads(settings.CONFIG_FILE.read_text(encoding='utf-8')))
    return render(request, 'prospectos/otras_computadoras.html', _ctx(
        request, 'Usar desde otras computadoras', activa_ahora=settings.EN_RED, guardado=guardado,
        direcciones=red.direcciones(settings.CONFIG), huella=red.huella_ca(settings.DATA_DIR),
        puerto=red.puerto(settings.CONFIG), sistema=settings.SISTEMA))


def conectar(request):
    """Página pública (solo con el acceso por red activo): cómo instalar el certificado en Windows o Mac."""
    from django.http import Http404
    from .servicios import red
    if not settings.EN_RED:
        raise Http404
    agente = request.META.get('HTTP_USER_AGENT', '')
    so = 'mac' if 'Macintosh' in agente else 'movil' if ('iPhone' in agente or 'iPad' in agente or 'Android' in agente) else 'windows'
    return render(request, 'prospectos/conectar.html', {
        'so': so, 'huella': red.huella_ca(settings.DATA_DIR), 'equipo': red.nombre_equipo(),
        'admin': request.build_absolute_uri('/admin/')})


def certificado(request):
    """Certificado público de la autoridad local (no es secreto: es lo que cada computadora instala)."""
    from django.http import FileResponse, Http404
    ruta = settings.DATA_DIR / 'red' / 'ca.crt'
    if not settings.EN_RED or not ruta.exists():
        raise Http404
    return FileResponse(open(ruta, 'rb'), content_type='application/x-x509-ca-cert', as_attachment=True,
                        filename='Prospectos-UDC.crt')


@staff_member_required
def ubicar(request):
    """Ubicar en el mapa las preparatorias que escriben los alumnos (buscar dirección o hacer clic)."""
    _permiso(request, 'change_preparatoria')
    from .servicios import preparatorias
    if request.method == 'POST':
        prepa = get_object_or_404(Preparatoria, pk=request.POST.get('prepa'))
        try:
            preparatorias.ubicar(prepa, request.POST.get('lat'), request.POST.get('lng'), request.POST.get('municipio'), request.user)
        except (ValueError, ArithmeticError) as e:
            return JsonResponse({'ok': False, 'error': str(e) if isinstance(e, ValueError) else 'Coordenadas inválidas.'}, status=400)
        return JsonResponse({'ok': True, 'id': prepa.pk, 'municipio': prepa.municipio})
    pendientes = list(reportes.Metricas.sin_ubicar())
    elegida = None
    if request.GET.get('prepa'):
        elegida = Preparatoria.objects.filter(pk=request.GET['prepa']).first()
        if elegida and elegida not in pendientes:
            elegida.total = elegida.prospectos.count()
            pendientes.insert(0, elegida)
    ubicadas = [{'id': p.pk, 'nombre': p.nombre, 'lat': float(p.lat), 'lng': float(p.lng)}
                for p in Preparatoria.objects.exclude(lat=None).exclude(lng=None)]
    lista = [{'id': p.pk, 'nombre': p.nombre, 'municipio': p.municipio, 'total': p.total, 'por_revisar': p.por_revisar,
              'lat': float(p.lat) if p.lat is not None else None, 'lng': float(p.lng) if p.lng is not None else None}
             for p in pendientes]
    return render(request, 'prospectos/ubicar.html', _ctx(
        request, 'Ubicar preparatorias en el mapa', lista=lista, ubicadas=ubicadas, tiles=TILES(),
        elegida=elegida.pk if elegida else (lista[0]['id'] if lista else None)))


# --------------------------------------------------------------------------- reporte ejecutivo
@staff_member_required
def reporte(request):
    _permiso(request, 'reporte_ejecutivo')
    hoy = timezone.localdate()
    desde = _fecha(request.GET.get('desde'), hoy - timedelta(days=90))
    hasta = _fecha(request.GET.get('hasta'), hoy)
    m = reportes.Metricas(desde, hasta)
    d = {
        'total': m.total(), 'pide_info': m.pide_informacion(), 'indecisos': m.indecisos(), 'costo': m.costo_por_prospecto(),
        'embudo': m.embudo(), 'carreras': m.carreras(), 'tasa': m.tasa_respuesta(), 'promotores': m.por_promotor(),
        'planteles': m.por_plantel(), 'canales': m.canales(), 'prepas': m.por_preparatoria_plantel(),
        'carreras_c': m.carreras(8, 'cuautitlan'), 'carreras_t': m.carreras(8, 'tultepec'),
        'factores': m.frecuencia('factor_eleccion', Cuestionario.ALUMNO, 'factores_eleccion'),
        'conocer': m.frecuencia('quiere_conocer', Cuestionario.ALUMNO, 'quiere_conocer'),
        'modalidad': m.frecuencia('modalidad', Cuestionario.ALUMNO, 'modalidades'),
        'temas': m.temas_duda(), 'segmentos': m.por_segmento(), 'agregadas': m.preguntas_agregadas(),
    }
    base = max(1, d['embudo'][0]['total'] if d['embudo'] else 1)
    for e in d['embudo']:
        e['pct'] = round(100 * e['total'] / base, 1)
    if request.GET.get('excel'):
        return _reporte_excel(request, d, desde, hasta)
    tope = max([r['tasa'] for r in d['tasa']] or [1]) or 1
    for r in d['tasa']:
        r['tasa_barra'] = round(100 * r['tasa'] / tope)
    graficas = {
        'g_carreras': _grafica(d['carreras'], color='var(--color-primary-600)'),
        'g_segmentos': json.dumps({'labels': [x['etiqueta'] for x in d['segmentos']], 'datasets': [{
            'data': [x['total'] for x in d['segmentos']],
            'backgroundColor': [COLORES_SEGMENTO.get(x['clave'], '#64748b') for x in d['segmentos']], 'borderWidth': 0}]}),
        'o_horizontal': _opciones_horizontales(), 'o_dona': _opciones_dona(),
    }
    return render(request, 'prospectos/reporte.html', _ctx(request, 'Reporte ejecutivo', d=d, desde=desde, hasta=hasta, **graficas))


def _reporte_excel(request, d, desde, hasta):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    ws.append(['Reporte ejecutivo', f'{desde:%d/%m/%Y} a {hasta:%d/%m/%Y}'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in [('Aspirantes', d['total']), ('Pidieron información (%)', d['pide_info']),
                 ('Aún no deciden carrera (%)', d['indecisos']), ('Costo total de eventos', d['costo']['costo_total']),
                 ('Costo por prospecto', d['costo']['costo_por_prospecto'])]:
        ws.append([k, v])
    secciones = [
        ('Embudo', ['Etapa', 'Total', '%'], [[e['etapa'], e['total'], e['pct']] for e in d['embudo']]),
        ('Carreras', ['Carrera', 'Total'], [[r['etiqueta'], r['total']] for r in d['carreras']]),
        ('Preparatorias', ['Preparatoria', 'Prospectos', 'Estimados', 'Tasa %'],
         [[r['etiqueta'], r['total'], r['estimados'], r['tasa']] for r in d['tasa']]),
        ('Promotores', ['Promotor', 'Total'], [[r['etiqueta'], r['total']] for r in d['promotores']]),
        ('Segmentos', ['Segmento', 'Total'], [[r['etiqueta'], r['total']] for r in d['segmentos']]),
        ('Planteles', ['Plantel', 'Aspirantes'], [[r['etiqueta'], r['total']] for r in d['planteles']]),
        ('Prepas por plantel', ['Preparatoria', 'Total', 'Cuautitlán', 'Tultepec'],
         [[r['etiqueta'], r['total'], r['cuautitlan'], r['tultepec']] for r in d['prepas']]),
        ('Quieren información por', ['Canal', 'Total', '%'], [[r['etiqueta'], r['total'], r['pct']] for r in d['canales']]),
        ('Qué buscan', ['Lo más importante', 'Total'], [[r['etiqueta'], r['total']] for r in d['factores']]),
        ('Qué quieren conocer', ['Tema', 'Total'], [[r['etiqueta'], r['total']] for r in d['conocer']]),
        ('Modalidad', ['Modalidad', 'Total'], [[r['etiqueta'], r['total']] for r in d['modalidad']]),
        *[(f"P{i} {a['titulo']}"[:31], ['Respuesta', 'Total'], [[r['etiqueta'], r['total']] for r in a['filas']])
          for i, a in enumerate(d['agregadas'], 1)],
    ]
    for nombre, enc, filas in secciones:
        h = wb.create_sheet(nombre[:31])
        h.append(enc)
        for c in h[1]:
            c.font = Font(bold=True)
        for f in filas:
            h.append(f)
        h.column_dimensions['A'].width = 40
    buf = BytesIO()
    wb.save(buf)
    auditoria.registrar('exportar_reporte', None, {'desde': str(desde), 'hasta': str(hasta)}, request=request)
    r = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r['Content-Disposition'] = f'attachment; filename="reporte-{desde:%Y%m%d}-{hasta:%Y%m%d}.xlsx"'
    return r


# --------------------------------------------------------------------------- configuración de scoring
NOMBRES_PESOS = {
    'pidio_informacion': 'Pidió información (WhatsApp o correo)',
    'eligio_carrera': 'Ya eligió al menos una carrera',
    'pregunta_requisitos': 'Quiere conocer los requisitos de inscripción',
    'eligio_modalidad': 'Ya eligió modalidad (semanal, sabatino o dominical)',
    'pregunta_costos': 'Quiere conocer costos o becas',
    'contacto_completo': 'Dejó teléfono y correo (alumno o familia)',
    'familia_pidio_informacion': 'La familia pidió información (solo si contestó)',
    'ingresa_proximo_ciclo': 'La familia planea que entre el próximo ciclo (solo si contestó)',
}
NOMBRES_UMBRALES = {
    'listo_para_convertir_score': 'Prioridad mínima para «Listo para inscribirse»',
    'indeciso_desde_carreras': 'Carreras marcadas a partir de las cuales es «Aún no decide carrera»',
}


@staff_member_required
def configuracion(request):
    _permiso(request, 'configurar_scoring')
    pesos = inteligencia.pesos()
    umbrales = {k: inteligencia.umbral(k) for k in NOMBRES_UMBRALES}
    if request.method == 'POST' and request.POST.get('accion') == 'permisos':
        return _guardar_permisos(request)
    if request.method == 'POST':
        try:
            nuevos = {k: max(0, min(100, int(request.POST.get(f'peso_{k}', v)))) for k, v in pesos.items()}
            nuevos_u = {k: max(0, min(100, int(request.POST.get(f'umbral_{k}', v)))) for k, v in umbrales.items()}
        except ValueError:
            messages.error(request, 'Todos los valores deben ser números enteros.')
            return redirect(request.path)
        if not sum(nuevos.values()):
            messages.error(request, 'Al menos un peso debe ser mayor que cero.')
            return redirect(request.path)
        Configuracion.guardar('score.pesos', nuevos, request.user)
        for k, v in nuevos_u.items():
            Configuracion.guardar(f'segmentos.{k}', v, request.user)
        auditoria.registrar('configurar_scoring', None, {'pesos': nuevos, 'umbrales': nuevos_u}, request=request)
        messages.success(request, f'Guardado. Se recalcularon {inteligencia.recalcular_todos()} prospectos.')
        return redirect(request.path)
    total = sum(pesos.values()) or 1
    filas = [{'clave': k, 'nombre': NOMBRES_PESOS.get(k, k), 'valor': v, 'pct': round(100 * v / total)} for k, v in pesos.items()]
    return render(request, 'prospectos/configuracion.html', _ctx(
        request, 'Reglas de prioridad y de permiso', filas=filas,
        umbrales=[{'clave': k, 'nombre': NOMBRES_UMBRALES[k], 'valor': v} for k, v in umbrales.items()],
        aviso=Configuracion.obtener('permisos.aviso_en_form') or {}, mayoria=settings.PROSPECTOS['EDAD_MAYORIA']))


def _guardar_permisos(request):
    """Solo un administrador: es una decisión legal (consentimiento tácito de mayores de edad)."""
    if not request.user.is_superuser:
        raise PermissionDenied
    activo = request.POST.get('aviso_activo') == '1'
    url = (request.POST.get('aviso_url') or '').strip()
    if activo and (not url.startswith('http') or request.POST.get('confirmo') != '1'):
        messages.error(request, 'Para activarlo pega el enlace del aviso de privacidad y confirma que aparece en el Form.')
        return redirect(request.path)
    valor = {'activo': activo, 'url': url, 'por': request.user.get_username(), 'en': timezone.now().isoformat()}
    Configuracion.guardar('permisos.aviso_en_form', valor, request.user)
    auditoria.registrar('configurar_permisos', None, {'aviso_en_form': activo}, request=request)
    n = inteligencia.recalcular_todos()
    messages.success(request, f'Reglas de permiso guardadas. Se recalcularon {n} aspirantes.')
    return redirect(request.path)


# --------------------------------------------------------------------------- mapeo de formulario
@staff_member_required
def mapeo(request, pk):
    _permiso(request, 'change_formulario')
    f = get_object_or_404(Formulario, pk=pk)
    codigos = list(Pregunta.objects.filter(cuestionario=f.tipo).values_list('codigo', 'texto'))
    if request.method == 'POST':
        propuesta = [{'question_id': qid, 'codigo': request.POST.get(f'q_{qid}') or None}
                     for qid in request.POST.getlist('qid')]
        usados = [p['codigo'] for p in propuesta if p['codigo']]
        repetidos = {c for c in usados if usados.count(c) > 1} - catalogos.REPETIBLES
        if repetidos:
            messages.error(request, 'Un código no puede asignarse a dos preguntas: ' + ', '.join(sorted(repetidos)))
        else:
            importacion.aplicar_mapeo(f, propuesta)
            n = importacion.reprocesar(f)
            auditoria.registrar('mapeo_formulario', f, {'preguntas': len(usados)}, request=request)
            messages.success(request, f'Mapeo guardado ({len(usados)} preguntas).'
                             + (f' Se reacomodaron {n} respuestas ya importadas.' if n else ''))
            return redirect(reverse('admin:prospectos_formulario_changelist'))
    try:
        propuesta = importacion.proponer_mapeo(f)
        error = None
    except Exception as e:  # noqa: BLE001 - sin acceso al Form: se muestra el error, no se rompe la página
        propuesta, error = [{'question_id': q, 'titulo': '(sin conexión con el Form)', 'codigo': c, 'similitud': None}
                            for q, c in (f.mapeo_campos or {}).items()], str(e)
    sin = set(dict(codigos)) - {p['codigo'] for p in propuesta if p['codigo']}
    return render(request, 'prospectos/mapeo.html', _ctx(request, f'Mapeo · {f}', formulario=f, propuesta=propuesta,
                                                           codigos=codigos, faltan=sorted(sin - catalogos.RECOMENDADAS),
                                                           recomendadas=sorted(sin & catalogos.RECOMENDADAS), error=error))


# --------------------------------------------------------------------------- respuestas de los Forms
def _respuestas_visibles(request):
    from .models import Respuesta
    qs = Respuesta.objects.select_related('prospecto__preparatoria', 'prospecto__carrera', 'evento__preparatoria', 'formulario')
    u = request.user
    if not u.is_superuser and u.has_perm('prospectos.ver_propios'):
        qs = qs.filter(Q(promotor=u) | Q(prospecto__promotor=u))
    return qs


@staff_member_required
def respuestas(request):
    """Todas las respuestas que llegan de los Forms, con su estado explicado en palabras simples."""
    from .models import EstadoRespuesta
    _permiso(request, 'view_prospecto')
    qs = _respuestas_visibles(request).prefetch_related('items').order_by('-enviada_en')
    f = request.GET
    if f.get('tipo') in (Cuestionario.ALUMNO, Cuestionario.TUTOR):
        qs = qs.filter(cuestionario=f['tipo'])
    if f.get('estado') in EstadoRespuesta.values:
        qs = qs.filter(estado=f['estado'])
    if f.get('q'):
        qs = qs.filter(Q(nombre_aspirante__icontains=f['q']) | Q(prospecto__nombre_aspirante__icontains=f['q']))
    conteos = dict(_respuestas_visibles(request).values_list('estado').annotate(n=Count('id')))
    from django.core.paginator import Paginator
    pagina = Paginator(qs, 40).get_page(request.GET.get('pagina'))
    from .admin import iniciales
    filas = []
    for r in pagina:
        etiqueta, tono, ayuda = reportes.explicar_estado(r)
        nombre = r.nombre_aspirante or (r.prospecto.nombre_aspirante if r.prospecto else 'Sin nombre')
        prepa = (r.prospecto.preparatoria if r.prospecto else None) or (r.evento.preparatoria if r.evento else None)
        partes = [str(prepa) if prepa else 'Sin preparatoria']
        if r.prospecto and r.prospecto.carrera:
            partes.append(str(r.prospecto.carrera))
        if r.evento:
            partes.append(f'evento {r.evento.codigo}')
        filas.append({'r': r, 'nombre': nombre, 'iniciales': iniciales(nombre), 'estado': etiqueta, 'tono': tono, 'ayuda': ayuda,
                      'detalle': ' · '.join(partes)})
    estados = [(k, *reportes.ESTADOS_LEGIBLES[k][:2], conteos.get(k, 0)) for k in
               (EstadoRespuesta.VALIDA, EstadoRespuesta.CUARENTENA, EstadoRespuesta.SIN_CONSENTIMIENTO, EstadoRespuesta.DESCARTADA)]
    return render(request, 'prospectos/respuestas.html', _ctx(request, 'Respuestas de los formularios', filas=filas, pagina=pagina,
                                                                estados=estados, f=f, total=sum(conteos.values())))


@staff_member_required
def respuesta(request, pk):
    _permiso(request, 'view_prospecto')
    r = get_object_or_404(_respuestas_visibles(request).prefetch_related('items'), pk=pk)
    from .models import EstadoRespuesta
    puede_permiso = (r.estado == EstadoRespuesta.CUARENTENA and r.prospecto_id
                     and request.user.has_perm('prospectos.change_prospecto'))
    if request.method == 'POST' and request.POST.get('accion') == 'permiso':
        if not puede_permiso:
            raise PermissionDenied
        if request.POST.get('confirmo') != '1':
            messages.error(request, 'Marca la casilla de confirmación: el tutor debe haber aceptado el aviso de privacidad.')
        else:
            try:
                inteligencia.registrar_permiso_tutor(r.prospecto, request.user, request.POST.get('tutor_nombre', ''),
                                                     request.POST.get('medio', ''), request.POST.get('nota', ''))
                messages.success(request, 'Permiso registrado. La respuesta ya se puede usar.')
            except ValueError as e:
                messages.error(request, str(e))
        return redirect(request.path)
    contacto = None
    if puede_permiso and request.GET.get('contacto') and request.user.has_perm('prospectos.ver_datos_personales'):
        p = r.prospecto
        contacto = [('Teléfono del aspirante', p.telefono), ('Teléfono del tutor', p.tutor_telefono), ('Tutor', p.tutor_nombre)]
        auditoria.registrar('ver_contacto_para_permiso', p, request=request)
    auditoria.registrar('ver_respuesta', r, request=request)
    etiqueta, tono, ayuda = reportes.explicar_estado(r)
    es_alumno = r.cuestionario == Cuestionario.ALUMNO
    return render(request, 'prospectos/respuesta.html', _ctx(
        request, f'Respuesta de {r.nombre_aspirante or "aspirante"}', r=r, temas=reportes.detalle_respuesta(r),
        estado=etiqueta, tono=tono, ayuda=ayuda, es_alumno=es_alumno, es_tutor=not es_alumno,
        puede_permiso=puede_permiso, contacto=contacto, medios=inteligencia.MEDIOS_PERMISO.items(),
        nombre=r.nombre_aspirante or (r.prospecto.nombre_aspirante if r.prospecto else 'Sin nombre')))
