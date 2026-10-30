"""Panel de administración (Django admin) con permisos por rol, enmascarado y auditoría."""

from django.contrib import admin, messages
from django.contrib.admin.options import ActionLocation
from django.contrib.auth.admin import GroupAdmin, UserAdmin
from django.contrib.auth.models import Group, User
from django.db import models
from django.db.models import Count, Q
from django.http import HttpResponseRedirect
from django.urls import reverse
from django.utils import timezone
from django.utils.html import format_html, format_html_join
from django.utils.safestring import mark_safe

from unfold.admin import ModelAdmin, StackedInline, TabularInline
from unfold.decorators import display
from unfold.forms import AdminPasswordChangeForm, UserChangeForm
from unfold.widgets import UnfoldAdminTextareaWidget

from .formularios_panel import CarreraForm, NuevoUsuarioForm, OpcionesFormSet, PreguntaForm, PreparatoriaForm
from .models import (etiqueta_segmento, Auditoria, Carrera, Configuracion, Cuestionario, Etiqueta, Evento, Formulario, Importacion, OpcionPregunta, Perfil,
                     PosibleCoincidencia, Pregunta, Preparatoria, Prospecto, Respaldo, Segmento, Seguimiento, TipoSeguimiento)
from .servicios import auditoria, familias, importacion, inteligencia, reportes, respaldos, seguimiento
from .servicios.cifrado import enmascarar_correo, enmascarar_telefono

admin.site.index_title = 'Inicio'


# --------------------------------------------------------------------------- usuarios
class AccionesFila:
    """Botones de editar (o ver) y eliminar al final de cada fila, solo los que el rol permite."""

    def get_list_display(self, request):
        base = list(super().get_list_display(request))
        admin_ = self
        info = (self.model._meta.app_label, self.model._meta.model_name)

        def acciones(obj):
            botones = []
            if admin_.has_change_permission(request, obj):
                botones.append(('admin:%s_%s_change' % info, 'edit', 'Editar', ''))
            elif admin_.has_view_permission(request, obj):
                botones.append(('admin:%s_%s_change' % info, 'visibility', 'Ver', ''))
            if admin_.has_delete_permission(request, obj):
                botones.append(('admin:%s_%s_delete' % info, 'delete', 'Eliminar', ' peligro'))
            return format_html('<span class="udc-fila-acciones">{}</span>', format_html_join(
                '', '<a class="udc-fila-boton{}" href="{}" title="{}" aria-label="{}"><span class="material-symbols-outlined">{}</span></a>',
                ((clase, reverse(url, args=[obj.pk]), texto, texto, icono) for url, icono, texto, clase in botones)))
        acciones.short_description = ''
        return base + [acciones]


class PerfilInline(StackedInline):
    model = Perfil
    can_delete = False
    verbose_name_plural = 'perfil'
    fields = ['codigo_promotor', 'debe_cambiar_clave']


admin.site.unregister(User)
admin.site.unregister(Group)


@admin.register(Group)
class RolAdmin(AccionesFila, GroupAdmin, ModelAdmin):
    """Solo administradores: quien edita roles puede darse a sí mismo cualquier permiso."""

    def has_module_permission(self, request):
        return request.user.is_superuser

    def has_view_permission(self, request, obj=None):
        return request.user.is_superuser

    def has_add_permission(self, request):
        return request.user.is_superuser

    def has_change_permission(self, request, obj=None):
        return request.user.is_superuser

    def has_delete_permission(self, request, obj=None):
        return request.user.is_superuser and obj is not None and obj.name not in ROLES_BASE


ROLES_BASE = {'Administrador', 'Admisiones', 'Coordinación', 'Promotor', 'Directivo'}


@admin.register(User)
class UsuarioAdmin(AccionesFila, UserAdmin, ModelAdmin):
    form = UserChangeForm
    add_form = NuevoUsuarioForm
    change_password_form = AdminPasswordChangeForm
    inlines = [PerfilInline]
    list_display = ['nombre', 'username', 'roles', 'dos_pasos', 'activo', 'last_login']
    list_filter = ['groups', 'is_active', 'perfil__dos_pasos_activo']
    search_fields = ['username', 'first_name', 'last_name', 'email']
    ordering = ['-is_active', 'first_name']
    actions = ['desactivar', 'reactivar', 'quitar_dos_pasos']
    add_fieldsets = ((None, {'fields': ('email', 'first_name', 'last_name', 'rol', 'codigo_promotor', 'clave'),
                             'description': 'Se crea en un paso. La contraseña temporal se muestra una sola vez al guardar.'}),)

    def get_fieldsets(self, request, obj=None):
        if not obj:
            return self.add_fieldsets
        campos = ['first_name', 'last_name', 'email', 'is_active', 'groups']
        if request.user.is_superuser:
            campos += ['is_superuser']
        return ((None, {'fields': ['username', 'password']}), ('Datos', {'fields': campos}),
                ('Fechas', {'fields': ['last_login', 'date_joined']}))

    def get_readonly_fields(self, request, obj=None):
        return ['last_login', 'date_joined'] if obj else []

    def get_inlines(self, request, obj):
        return [PerfilInline] if obj else []  # al crear, el formulario ya arma el perfil

    def get_queryset(self, request):
        qs = super().get_queryset(request).prefetch_related('groups').select_related('perfil')
        return qs if request.user.is_superuser else qs.filter(is_superuser=False)  # nadie toca a un administrador salvo otro

    def formfield_for_manytomany(self, db_field, request, **kwargs):
        if db_field.name == 'groups' and not request.user.is_superuser:
            kwargs['queryset'] = Group.objects.exclude(name='Administrador')
        return super().formfield_for_manytomany(db_field, request, **kwargs)

    def get_form(self, request, obj=None, **kwargs):
        form = super().get_form(request, obj, **kwargs)
        if obj is None and not request.user.is_superuser:
            class Restringido(form):
                def __init__(self, *a, **k):
                    super().__init__(*a, **k)
                    self.fields['rol'].queryset = Group.objects.exclude(name='Administrador').order_by('name')
            return Restringido
        return form

    def has_delete_permission(self, request, obj=None):
        return False  # se desactiva: borrar dejaría la bitácora y los seguimientos sin autor

    @display(description='Nombre', ordering='first_name')
    def nombre(self, obj):
        return obj.get_full_name() or obj.username

    @display(description='Roles')
    def roles(self, obj):
        return ', '.join(g.name for g in obj.groups.all()) or ('Administrador' if obj.is_superuser else '—')

    @display(description='Activo', boolean=True, ordering='is_active')
    def activo(self, obj):
        return obj.is_active

    @display(description='2 pasos', boolean=True)
    def dos_pasos(self, obj):
        try:
            return obj.perfil.dos_pasos_activo
        except Perfil.DoesNotExist:
            return False

    @admin.action(description='Quitar la verificación en 2 pasos (perdió su teléfono y sus códigos)')
    def quitar_dos_pasos(self, request, queryset):
        if not request.user.is_superuser:
            self.message_user(request, 'Solo un administrador puede hacerlo.', messages.ERROR)
            return
        from .servicios import dos_pasos
        n = 0
        for u in queryset.exclude(pk=request.user.pk):
            dos_pasos.desactivar(Perfil.de(u))
            auditoria.registrar('dos_pasos_quitar', u, {'por': request.user.get_username()}, request=request)
            n += 1
        self.message_user(request, f'Se quitó la verificación en 2 pasos a {n} usuario(s). Los administradores tendrán que '
                                   'configurarla de nuevo al entrar.' + (' La tuya no se puede quitar desde aquí: usa Mi perfil.'
                                                                          if queryset.filter(pk=request.user.pk).exists() else ''))

    @admin.action(description='Desactivar (ya no puede entrar; se conserva su historial)')
    def desactivar(self, request, queryset):
        n = queryset.exclude(pk=request.user.pk).update(is_active=False)
        auditoria.registrar('usuario_desactivar', None, {'usuarios': n}, request=request)
        self.message_user(request, f'{n} usuario(s) desactivado(s).' + (' No puedes desactivarte a ti mismo.'
                                                                          if queryset.filter(pk=request.user.pk).exists() else ''))

    @admin.action(description='Reactivar')
    def reactivar(self, request, queryset):
        n = queryset.update(is_active=True)
        self.message_user(request, f'{n} usuario(s) reactivado(s).')

    def save_model(self, request, obj, form, change):
        super().save_model(request, obj, form, change)
        auditoria.registrar('usuario_' + ('editar' if change else 'crear'), obj, {'campos': list(form.changed_data)}, request=request)
        if not change and getattr(form, 'clave_generada', None):
            self.message_user(request, format_html(
                'Usuario creado. Contraseña temporal de <b>{}</b>: <code style="font-size:1.1em">{}</code> — '
                'dásela en persona o por un medio privado; se le pedirá cambiarla al entrar. No se vuelve a mostrar.',
                obj.username, form.clave_generada), messages.WARNING)

    def user_change_password(self, request, id, form_url=''):
        respuesta = super().user_change_password(request, id, form_url)
        if request.method == 'POST' and respuesta.status_code == 302 and str(id) != str(request.user.pk):
            Perfil.objects.update_or_create(usuario_id=id, defaults={'debe_cambiar_clave': True})
            u = User.objects.get(pk=id)
            from django.core.cache import cache
            cache.delete(f'acceso:u:{u.get_username().strip().lower()[:150]}')  # quita el bloqueo por intentos
            auditoria.registrar('usuario_restablecer_clave', u, None, request=request)
        return respuesta


# --------------------------------------------------------------------------- catálogos
@admin.register(Preparatoria)
class PreparatoriaAdmin(AccionesFila, ModelAdmin):
    list_display = ['nombre', 'municipio', 'ubicada', 'por_revisar', 'prospectos', 'alumnos_estimados', 'qr', 'activo']
    list_filter = ['por_revisar', 'activo', 'municipio']
    search_fields = ['nombre', 'municipio', 'alias']
    actions = ['unir_preparatorias']
    form = PreparatoriaForm
    fields = ['nombre', 'municipio', 'lat', 'lng', 'alias', 'alumnos_estimados', 'por_revisar', 'activo']

    def get_queryset(self, request):
        return super().get_queryset(request).annotate(
            n=Count('prospectos', filter=Q(prospectos__fusionado_en__isnull=True)))

    @display(description='Aspirantes', ordering='n')
    def prospectos(self, obj):
        return obj.n

    @display(description='QR permanente')
    def qr(self, obj):
        return format_html('<a class="udc-mini" href="{}"><span class="material-symbols-outlined">qr_code_2</span>{}</a>',
                           reverse('prospectos:prepa_qr', args=[obj.pk]), obj.codigo_qr)

    @display(description='En el mapa')
    def ubicada(self, obj):
        if obj.lat is not None and obj.lng is not None:
            return mark_safe('<span class="udc-chip udc-seg-listo_para_convertir">Sí</span>')  # nosec B308
        return format_html('<a class="udc-chip udc-seg-sensible_beca" href="{}?prepa={}">Ubicar</a>',
                           reverse('prospectos:ubicar'), obj.pk)

    @admin.action(description='Unir: son la misma preparatoria escrita distinto')
    def unir_preparatorias(self, request, queryset):
        from .servicios import preparatorias
        try:
            queda = preparatorias.unir(list(queryset), request.user)
        except ValueError as e:
            self.message_user(request, str(e), messages.ERROR)
            return
        self.message_user(request, f'Unidas en «{queda}». Los nombres de las otras quedaron como «otras formas de escribirla»: '
                                   'las próximas respuestas que lo escriban así caerán aquí solas.', messages.SUCCESS)


class AvisaPublicar:
    """Tras cambiar preguntas, opciones o carreras: recuerda que el Google Form todavía no tiene el cambio."""

    def _tipos(self, obj):
        return [obj.cuestionario] if isinstance(obj, Pregunta) else ['ALUMNO', 'TUTOR']

    def _marcar(self, request, obj):
        for t in self._tipos(obj):
            Configuracion.guardar(f'publicar.pendiente.{t}', True, request.user)
        if Formulario.objects.filter(activo=True).exclude(google_form_id__startswith='demo-').exists():
            self.message_user(request, format_html(
                'El Google Form todavía no tiene este cambio. <a class="underline" href="{}">Pásalo al Form en Encuestas</a>.',
                reverse('prospectos:encuestas')), messages.INFO)

    def save_model(self, request, obj, form, change):
        super().save_model(request, obj, form, change)
        self._marcar(request, obj)

    def delete_model(self, request, obj):
        tipos = self._tipos(obj)
        super().delete_model(request, obj)
        for t in tipos:
            Configuracion.guardar(f'publicar.pendiente.{t}', True, request.user)


@admin.register(Carrera)
class CarreraAdmin(AccionesFila, AvisaPublicar, ModelAdmin):
    form = CarreraForm
    list_display = ['nombre', 'area', 'planteles_txt', 'interesados_n', 'activo']
    list_editable = ['activo']

    def get_queryset(self, request):
        return super().get_queryset(request).annotate(n=Count('interesados', filter=Q(interesados__fusionado_en__isnull=True)))

    @display(description='Aspirantes', ordering='n')
    def interesados_n(self, obj):
        return obj.n

    def has_delete_permission(self, request, obj=None):
        # Se apaga en vez de borrar: borrar dejaría a los aspirantes sin la carrera que eligieron.
        return super().has_delete_permission(request, obj) and (obj is None or not obj.interesados.exists())

    def save_model(self, request, obj, form, change):
        obj.editada = True
        super().save_model(request, obj, form, change)
        auditoria.registrar('carrera_' + ('editar' if change else 'crear'), obj, {'campos': list(form.changed_data)}, request=request)
        if not change or 'alias' in form.changed_data or 'nombre' in form.changed_data:
            n = _releer(Cuestionario.ALUMNO) + _releer(Cuestionario.TUTOR)
            if n:
                self.message_user(request, f'Se volvieron a leer las respuestas: {n} cambiaron con esta carrera.')

    @display(description='Planteles')
    def planteles_txt(self, obj):
        from .catalogos import PLANTELES_CORTO
        return ', '.join(PLANTELES_CORTO.get(p, p) for p in obj.planteles or []) or '—'

    list_filter = ['area', 'activo']
    search_fields = ['nombre']


@admin.register(Etiqueta)
class EtiquetaAdmin(AccionesFila, ModelAdmin):
    search_fields = ['nombre']


def _releer(cuestionario) -> int:
    """Vuelve a leer las respuestas guardadas de los Forms de ese tipo (sin pedir nada a Google)."""
    return sum(importacion.reprocesar(f) for f in Formulario.objects.filter(tipo=cuestionario, activo=True))


class OpcionInline(TabularInline):
    model = OpcionPregunta
    formfield_overrides = {models.TextField: {'widget': UnfoldAdminTextareaWidget(attrs={'rows': 1})}}
    formset = OpcionesFormSet
    fields = ['texto', 'alias', 'orden', 'fija_txt']
    readonly_fields = ['fija_txt']
    extra = 1
    tab = False

    @display(description='')
    def fija_txt(self, obj):
        if obj and obj.fija:
            return mark_safe('<span class="udc-chip udc-seg-sensible_beca" title="La usan las reglas de prioridad: '  # nosec B308
                             'puedes cambiarle el texto, no borrarla">Usada por las reglas</span>')
        return ''


@admin.register(Pregunta)
class PreguntaAdmin(AccionesFila, AvisaPublicar, ModelAdmin):
    form = PreguntaForm
    list_display = ['etiqueta_txt', 'cuestionario', 'tipo_txt', 'tema', 'opciones_n', 'origen', 'orden']
    list_display_links = ['etiqueta_txt']
    list_filter = ['cuestionario', 'tema', 'de_sistema', 'tipo']
    search_fields = ['codigo', 'texto', 'etiqueta']
    ordering = ['cuestionario', 'orden', 'id']
    inlines = [OpcionInline]

    def get_queryset(self, request):
        return super().get_queryset(request).annotate(n_op=Count('opciones_lista'))

    def get_inlines(self, request, obj):
        if obj is None:
            return [OpcionInline]  # al crear todavía no se sabe el tipo: se muestran por si es de opciones
        return [OpcionInline] if obj.con_opciones else []

    def get_fieldsets(self, request, obj=None):
        return [(None, {'fields': ['cuestionario', 'texto', 'etiqueta', 'tema', 'tipo', 'orden', 'codigo'],
                        'description': ('Para agregar una pregunta: créala aquí <b>y</b> en el Google Form, con el mismo texto. '
                                        'Luego en Formularios → «Sincronizar con Google». Las respuestas que ya llegaron se '
                                        'vuelven a leer solas.')})]

    @display(description='Pregunta', ordering='etiqueta')
    def etiqueta_txt(self, obj):
        return obj.etiqueta or obj.texto[:70]

    @display(description='Tipo')
    def tipo_txt(self, obj):
        from .models import TipoPregunta
        return dict(TipoPregunta.choices).get(obj.tipo) or {'catalogo': 'Catálogo del programa', 'telefono': 'Teléfono',
                                                             'correo': 'Correo', 'consentimiento': 'Sí / No (aviso)',
                                                             'oculto': 'Uso interno (QR)'}.get(obj.tipo, obj.tipo)

    @display(description='Opciones', ordering='n_op')
    def opciones_n(self, obj):
        return obj.n_op or '—'

    @display(description='Origen', label={'Del sistema': 'info', 'Agregada': 'success'})
    def origen(self, obj):
        return 'Del sistema' if obj.de_sistema else 'Agregada'

    def has_delete_permission(self, request, obj=None):
        return super().has_delete_permission(request, obj) and (obj is None or not obj.de_sistema)

    def save_model(self, request, obj, form, change):
        obj.editada = True
        if obj.tipo in ('telefono', 'correo'):
            obj.sensible = True
        super().save_model(request, obj, form, change)

    def save_related(self, request, form, formsets, change):
        super().save_related(request, form, formsets, change)
        obj = form.instance
        if not obj.con_opciones:
            obj.opciones_lista.filter(fija=False).delete()  # si cambió a texto/número, las opciones sobran
        cambio = form.changed_data or any(fs.has_changed() for fs in formsets)
        auditoria.registrar('pregunta_' + ('editar' if change else 'crear'), obj, {'campos': list(form.changed_data)}, request=request)
        if cambio:
            n = _releer(obj.cuestionario)
            self.message_user(request, f'Guardado. Se volvieron a leer las respuestas ({n} cambiaron).' if n else 'Guardado.')

    def delete_model(self, request, obj):
        auditoria.registrar('pregunta_borrar', obj, {'codigo': obj.codigo}, request=request)
        super().delete_model(request, obj)


# --------------------------------------------------------------------------- eventos
@admin.register(Evento)
class EventoAdmin(AccionesFila, ModelAdmin):
    list_display = ['preparatoria', 'fecha', 'estado_label', 'promotor', 'registros', 'herramientas']
    list_fullwidth = True
    list_filter = ['estado', 'fecha', 'promotor', 'preparatoria']
    search_fields = ['codigo', 'preparatoria__nombre']
    date_hierarchy = 'fecha'
    autocomplete_fields = ['preparatoria']
    readonly_fields = ['codigo']
    fieldsets = [
        (None, {'fields': ['preparatoria', 'fecha', ('hora_inicio', 'hora_fin'), 'actividad', 'estado', 'promotor', 'codigo']}),
        ('Para la solicitud a la preparatoria', {'fields': [('dirigido_a', 'dirigido_cargo'), 'grupos', 'alumnos_esperados', 'requerimientos']}),
        ('Cartas de beca para los alumnos', {'fields': ['beca_carta', ('beca_porcentaje', 'beca_hojas'), ('beca_vigencia', 'beca_inicio')],
                                             'description': 'Se imprimen desde «Becas» en la lista de eventos, una por alumno y con folio.'}),
        ('Resultados', {'fields': ['alumnos_asistentes', 'meta_registros', 'costo', 'notas']}),
    ]

    @display(description='Estado', ordering='estado', label={'Por solicitar': 'warning', 'Solicitado': 'info', 'Confirmado': 'primary',
                                                             'Realizado': 'success', 'Cancelado': 'danger'})
    def estado_label(self, obj):
        return obj.get_estado_display()

    def get_queryset(self, request):
        qs = super().get_queryset(request).select_related('preparatoria', 'promotor').annotate(
            n=Count('prospectos', filter=Q(prospectos__fusionado_en__isnull=True)))
        if solo_propios(request, 'ver_eventos_propios'):
            qs = qs.filter(promotor=request.user)
        return qs

    @display(description='Prospectos', ordering='n')
    def registros(self, obj):
        if obj.meta_registros:
            pct = min(100, round(100 * obj.n / obj.meta_registros))
            return format_html('<span class="udc-meta"><b>{}</b><span class="udc-barra"><i style="width:{}%"></i></span></span>', obj.n, pct)
        return obj.n

    @display(description='')
    def herramientas(self, obj):
        return format_html(
            '<span class="udc-acciones"><a class="udc-mini" href="{}"><span class="material-symbols-outlined">description</span>Solicitud</a>'
            '<a class="udc-mini" href="{}"><span class="material-symbols-outlined">qr_code_2</span>QR</a>'
            '<a class="udc-mini" href="{}"><span class="material-symbols-outlined">style</span>Tarjetas</a>'
            '<a class="udc-mini udc-mini-vivo" href="{}"><span class="material-symbols-outlined">sensors</span>En vivo</a>'
            '<a class="udc-mini" href="{}"><span class="material-symbols-outlined">summarize</span>Reporte</a>{}</span>',
            reverse('prospectos:evento_solicitud', args=[obj.pk]), reverse('prospectos:evento_qr', args=[obj.pk]),
            reverse('prospectos:tarjetas', args=[obj.pk]), reverse('prospectos:evento_vivo', args=[obj.pk]),
            reverse('prospectos:evento_reporte', args=[obj.pk]),
            format_html('<a class="udc-mini" href="{}"><span class="material-symbols-outlined">workspace_premium</span>Becas ({})</a>',
                        reverse('prospectos:evento_becas', args=[obj.pk]), obj.beca_hojas or 0) if obj.beca_carta else '')


def solo_propios(request, permiso: str) -> bool:
    u = request.user
    return not u.is_superuser and u.has_perm(f'prospectos.{permiso}')


# --------------------------------------------------------------------------- formularios
class ImportacionInline(TabularInline):
    model = Importacion
    extra = 0
    max_num = 0
    can_delete = False
    fields = readonly_fields = ['iniciada_en', 'terminada_en', 'estado', 'nuevas', 'actualizadas', 'omitidas', 'mensaje']

    def get_queryset(self, request):
        ids = Importacion.objects.order_by('-iniciada_en').values_list('id', flat=True)[:500]
        return super().get_queryset(request).filter(id__in=list(ids))


@admin.register(Formulario)
class FormularioAdmin(AccionesFila, ModelAdmin):
    list_display = ['nombre', 'tipo', 'activo', 'ultima_importacion_en', 'ultima_respuesta_en', 'estado', 'mapeo']
    readonly_fields = ['ultima_respuesta_en', 'ultima_importacion_en', 'ultimo_error', 'errores_consecutivos']
    actions = ['sincronizar_google', 'importar_ahora']
    fieldsets = [
        (None, {'fields': ['nombre', 'tipo', 'google_form_id', 'url_publica', 'activo']}),
        ('Campos prellenados (QR)', {'fields': ['campos_prellenado'],
                                     'description': 'Ejemplo: {"evento": "entry.123", "promotor": "entry.456", "token_familia": "entry.789"}'}),
        ('Mapeo (usa «Revisar mapeo» en la lista)', {'fields': ['mapeo_campos'], 'classes': ['collapse']}),
        ('Estado del importador', {'fields': readonly_fields}),
    ]

    def get_inlines(self, request, obj):
        return [ImportacionInline] if obj else []

    @display(description='Estado', label={'OK': 'success', 'Sin importar': 'info'})
    def estado(self, obj):
        if obj.ultimo_error:
            return format_html('<span class="udc-chip udc-seg-no_contactar" title="{}">Error ×{}</span>', obj.ultimo_error,
                               obj.errores_consecutivos)
        return 'OK' if obj.ultima_importacion_en else 'Sin importar'

    @display(description='')
    def mapeo(self, obj):
        return format_html('<a class="udc-mini" href="{}"><span class="material-symbols-outlined">account_tree</span>Revisar mapeo ({})</a>',
                           reverse('prospectos:mapeo', args=[obj.pk]), len(obj.mapeo_campos or {}))

    @admin.action(description='Sincronizar con Google (URL pública, mapeo y campos del QR)')
    def sincronizar_google(self, request, queryset):
        from .servicios.fuentes import ErrorFuente
        for f in queryset:
            try:
                r = importacion.sincronizar(f)
                faltan = f" · sin pregunta: {', '.join(r['faltan'])}" if r['faltan'] else ''
                faltan += f" · respuestas reacomodadas: {r['reprocesadas']}" if r.get('reprocesadas') else ''
                dudosos = [c for c, o in (r.get('prellenado') or {}).items() if o == 'deducido']
                aviso = (f" · Comprueba con «Obtener enlace prellenado» los entry de: {', '.join(dudosos)}" if dudosos else '')
                self.message_user(request, f"{f}: {r['mapeadas']} de {r['preguntas']} preguntas mapeadas{faltan}{aviso}. "
                                  'Revisa el mapeo antes de importar.', messages.WARNING if (r['faltan'] or dudosos) else messages.SUCCESS)
            except ErrorFuente as e:
                self.message_user(request, f'{f}: {e}', messages.ERROR)

    @admin.action(description='Importar ahora')
    def importar_ahora(self, request, queryset):
        for f in queryset:
            r = importacion.importar(f)
            nivel = messages.SUCCESS if r.estado == 'ok' else messages.ERROR
            self.message_user(request, f'{f}: {r.estado} · nuevas {r.nuevas} · actualizadas {r.actualizadas}'
                              + (f' · {r.mensaje}' if r.mensaje else ''), nivel)


# --------------------------------------------------------------------------- prospectos
class VistaFilter(admin.SimpleListFilter):
    title = 'vista'
    parameter_name = 'vista'

    def lookups(self, request, model_admin):
        return [('todos', 'Todos (incluye sin permiso)'), ('cuarentena', 'Esperando permiso'),
                ('fusionados', 'Fusionados'), ('anonimizados', 'Anonimizados')]

    def choices(self, changelist):
        yield {'selected': self.value() is None, 'query_string': changelist.get_query_string(remove=[self.parameter_name]),
               'display': 'Con permiso (utilizables)'}
        yield from list(super().choices(changelist))[1:]

    def queryset(self, request, qs):
        v = self.value()
        if v == 'todos':
            return qs.filter(fusionado_en__isnull=True)
        if v == 'cuarentena':
            return qs.filter(consentimiento=False, fusionado_en__isnull=True, anonimizado_en__isnull=True)
        if v == 'fusionados':
            return qs.filter(fusionado_en__isnull=False)
        if v == 'anonimizados':
            return qs.filter(anonimizado_en__isnull=False)
        return qs.utilizables()


class SegmentoFilter(admin.SimpleListFilter):
    title = 'grupo'
    parameter_name = 'segmento'

    def lookups(self, request, model_admin):
        return Segmento.choices

    def queryset(self, request, qs):
        if self.value():
            return qs.filter(segmentos__icontains=f'"{self.value()}"')
        return qs


class ScoreFilter(admin.SimpleListFilter):
    title = 'prioridad'
    parameter_name = 'score'

    def lookups(self, request, model_admin):
        return [('alto', '70 o más'), ('medio', '40 a 69'), ('bajo', 'Menos de 40')]

    def queryset(self, request, qs):
        return {'alto': qs.filter(score__gte=70), 'medio': qs.filter(score__gte=40, score__lt=70),
                'bajo': qs.filter(score__lt=40)}.get(self.value(), qs)


class NuevoSeguimientoInline(StackedInline):
    """Registrar un contacto o nota nuevos."""
    model = Seguimiento
    verbose_name = 'nuevo contacto o nota'
    verbose_name_plural = 'registrar contacto o nota'
    extra = 1
    fields = [('tipo', 'fecha'), 'notas', 'vence_en']
    can_delete = False
    show_title = False

    def get_queryset(self, request):
        return super().get_queryset(request).none()

    def formfield_for_choice_field(self, db_field, request, **kwargs):
        if db_field.name == 'tipo':
            kwargs['choices'] = [c for c in TipoSeguimiento.choices
                                 if c[0] not in (TipoSeguimiento.CAMBIO_ESTATUS, TipoSeguimiento.SISTEMA)]
        return super().formfield_for_choice_field(db_field, request, **kwargs)


def nivel_score(score: int) -> str:
    return 'alto' if score >= 70 else 'medio' if score >= 40 else 'bajo'


def pill_score(score: int):
    return format_html('<span class="udc-score udc-score-{}">{}</span>', nivel_score(score), score)


def chips_segmentos(segmentos, maximo=None):
    if not segmentos:
        return mark_safe('<span class="text-subtle">—</span>')  # noqa: S308 - texto fijo  # nosec B308
    visibles = segmentos[:maximo] if maximo else segmentos
    resto = len(segmentos) - len(visibles)
    return format_html('<span class="udc-chips">{}{}</span>', format_html_join(
        '', '<span class="udc-chip udc-seg-{}">{}</span>', ((s, etiqueta_segmento(s)) for s in visibles)),
        format_html('<span class="udc-chip udc-seg-mas" title="{}">+{}</span>',
                    ', '.join(etiqueta_segmento(s) for s in segmentos[len(visibles):]), resto) if resto else '')


ESTATUS_COLOR = {'Registrado': 'info', 'Contactado': 'primary', 'Asistió a open house': 'warning',
                 'Solicitó inscripción': 'warning', 'Inscrito': 'success', 'Descartado': 'danger'}


def iniciales(nombre: str) -> str:
    partes = [p for p in (nombre or '').split() if p[:1].isalpha()]
    return ''.join(p[0] for p in partes[:2]).upper() or '·'


@admin.register(Prospecto)
class ProspectoAdmin(AccionesFila, ModelAdmin):
    change_form_template = 'admin/prospectos/prospecto/change_form.html'
    list_display = ['aspirante', 'score_badge', 'familia', 'segmentos_txt', 'plantel', 'carrera', 'canal', 'estatus_label',
                    'ultimo_contacto_en', 'asesor']
    list_filter_submit = True
    list_fullwidth = True
    list_filter = [VistaFilter, 'plantel', 'estatus', SegmentoFilter, ScoreFilter, 'canal_informacion', 'no_contactar',
                   'preparatoria', 'carreras', 'evento', 'asesor', 'promotor', 'etiquetas', 'tiene_tutor']
    search_fields = ['nombre_aspirante', 'nombre_normalizado', 'tutor_nombre', 'token_familia']
    list_select_related = ['preparatoria', 'carrera', 'asesor']
    list_per_page = 50
    inlines = [NuevoSeguimientoInline]
    actions = ['exportar_excel', 'exportar_excel_contacto', 'asignarme', 'recalcular', 'anonimizar_arco']
    autocomplete_fields = ['etiquetas']
    fieldsets = [
        ('Seguimiento', {'fields': ['estatus', 'asesor', 'etiquetas']}),
        ('Aspirante', {'classes': ['tab'], 'fields': ['nombre_aspirante', 'edad', 'preparatoria', 'plantel', 'carreras_txt', 'modalidad',
                                  'canal_informacion', 'no_contactar', 'evento', 'promotor', 'tutor_nombre']}),
        ('Prioridad y grupos', {'classes': ['tab'], 'fields': ['score', 'detalle_score', 'segmentos_txt', 'brechas_txt']}),
        ('Permiso', {'classes': ['tab'], 'fields': ['consentimiento', 'origen_txt', 'consentimiento_en', 'permiso_registrado_por',
                                'permiso_nota', 'consentimiento_negado', 'bienvenida_enviada_en', 'anonimizado_en', 'fusionado_en', 'creado_en']}),
    ]

    def get_readonly_fields(self, request, obj=None):
        editables = {'estatus', 'asesor', 'etiquetas'}
        return [f for fs in self.fieldsets for f in fs[1]['fields'] if f not in editables]

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        if solo_propios(request, 'ver_propios'):
            qs = qs.filter(promotor=request.user)
        return qs

    def has_add_permission(self, request):
        return False  # los prospectos solo nacen del importador

    def has_delete_permission(self, request, obj=None):
        return False  # se anonimiza, no se borra (ver retención)

    # --- columnas
    @display(description='Aspirante', header=True, ordering='nombre_aspirante')
    def aspirante(self, obj):
        return [obj.nombre_aspirante, obj.preparatoria.nombre if obj.preparatoria else 'Sin preparatoria', iniciales(obj.nombre_aspirante)]

    @display(description='Quién contestó')
    def familia(self, obj):
        if obj.tiene_alumno and obj.tiene_tutor:
            return format_html('Alumno y familia<br><small class="text-subtle">{}</small>', obj.tutor_nombre or '')
        return 'Solo el alumno' if obj.tiene_alumno else format_html('Solo la familia<br><small class="text-subtle">{}</small>', obj.tutor_nombre or '')

    @display(description='Prioridad', ordering='score')
    def score_badge(self, obj):
        return pill_score(obj.score)

    @display(description='Grupos')
    def segmentos_txt(self, obj):
        return chips_segmentos(obj.segmentos, 2)

    @display(description='Estatus', ordering='estatus', label=ESTATUS_COLOR)
    def estatus_label(self, obj):
        return obj.get_estatus_display()

    @display(description='Alumno + tutor', boolean=True)
    def pareja(self, obj):
        return obj.tiene_alumno and obj.tiene_tutor

    @display(description='Prefiere', ordering='canal_informacion', label={'WhatsApp': 'success', 'Correo': 'info', 'Familia: sí': 'info', 'No quiere': 'danger'})
    def canal(self, obj):
        return {'whatsapp': 'WhatsApp', 'correo': 'Correo', 'si': 'Familia: sí', 'no': 'No quiere'}.get(obj.canal_informacion)

    @display(description='Carreras de interés')
    def carreras_txt(self, obj):
        return ', '.join(c.nombre for c in obj.carreras.all()) or 'Aún no decide'

    @display(description='Cómo se obtuvo')
    def origen_txt(self, obj):
        return {'tutor': 'Cuestionario del tutor', 'asesor': 'Permiso del tutor registrado por un asesor',
                'alumno': 'Mayor de edad: aceptó el aviso en el Form',
                'aviso': 'Mayor de edad: el Form muestra el aviso (consentimiento tácito)'}.get(obj.origen_consentimiento, '—')

    @display(description='Diferencias', ordering='tiene_brecha', label={'Sí': 'danger'})
    def brecha(self, obj):
        return 'Sí' if obj.tiene_brecha else None

    @display(description='De dónde salen los puntos')
    def detalle_score(self, obj):
        from .servicios.inteligencia import pesos
        from .views import NOMBRES_PESOS as nombres
        filas = [('si' if k in (obj.score_detalle or {}) else 'no', nombres.get(k, k), f'+{v}' if k in (obj.score_detalle or {}) else '—')
                 for k, v in ((k, (obj.score_detalle or {}).get(k, 0)) for k in pesos())]
        return format_html('<div class="udc-score-grande">{}<span>de 100</span></div><ul class="udc-reglas">{}</ul>', pill_score(obj.score),
                           format_html_join('', '<li class="udc-regla-{}"><span>{}</span><b>{}</b></li>', filas))

    @display(description='Diferencias alumno / tutor')
    def brechas_txt(self, obj):
        if not obj.brechas:
            return '—'
        return format_html_join('', '<div class="udc-brecha"><span class="material-symbols-outlined">warning</span><div><b>{}</b>'
                                '<small>Alumno: {} · Tutor: {}</small></div></div>',
                                ((b['mensaje'], b.get('alumno') or '—', b.get('tutor') or '—') for b in obj.brechas))

    # --- detalle
    def change_view(self, request, object_id, form_url='', extra_context=None):
        obj = self.get_object(request, object_id)
        extra = extra_context or {}
        if obj:
            auditoria.registrar('ver', obj, request=request)
            puede = request.user.has_perm('prospectos.ver_datos_personales')
            revelar = puede and request.GET.get('contacto') == '1'
            if revelar:
                auditoria.registrar('ver_datos_personales', obj, {'campos': Prospecto.CAMPOS_CONTACTO}, request=request)
            extra.update({
                'iniciales': iniciales(obj.nombre_aspirante),
                'segmentos_detalle': [(x, etiqueta_segmento(x)) for x in obj.segmentos or []],
                'comparativo': reportes.comparativo(obj),
                'historial': obj.seguimientos.select_related('usuario').order_by('-fecha')[:100],
                'puede_ver_contacto': puede, 'contacto_revelado': revelar,
                'contacto': [
                    ('Correo del alumno', obj.correo if revelar else enmascarar_correo(obj.correo)),
                    ('Teléfono del alumno', obj.telefono if revelar else enmascarar_telefono(obj.telefono)),
                    ('Correo del tutor', obj.tutor_correo if revelar else enmascarar_correo(obj.tutor_correo)),
                    ('Teléfono del tutor', obj.tutor_telefono if revelar else enmascarar_telefono(obj.tutor_telefono)),
                ] + ([('CURP', obj.curp if revelar else obj.curp[:4] + '••••••••••••••')] if obj.curp else []),
                'familia_posible': familias.candidatos(obj, maximo=2) if obj.tiene_alumno != obj.tiene_tutor else [],
                'absorbidos': list(obj.absorbidos.filter(anonimizado_en__isnull=True).only('pk', 'nombre_aspirante', 'tutor_nombre')),
                'coincidencias': PosibleCoincidencia.objects.filter(Q(prospecto=obj) | Q(candidato=obj), estado='pendiente')
                .exclude(motivo__startswith=familias.MOTIVO)
                .select_related('prospecto', 'candidato'),
            })
        return super().change_view(request, object_id, form_url, extra)

    def save_model(self, request, obj, form, change):
        if 'estatus' in form.changed_data:
            nuevo, obj.estatus = obj.estatus, form.initial['estatus']
            super().save_model(request, obj, form, change)
            seguimiento.cambiar_estatus(obj, nuevo, request.user)
        else:
            super().save_model(request, obj, form, change)
        if form.changed_data:
            auditoria.registrar('editar', obj, {'campos': list(form.changed_data)}, request=request)
        if 'asesor' in form.changed_data and obj.asesor_id:
            from .servicios import avisos
            avisos.asignado(obj.asesor, obj, por=request.user)

    def save_formset(self, request, form, formset, change):
        formset.new_objects, formset.changed_objects, formset.deleted_objects = [], [], []
        for f in formset.forms:
            if not f.has_changed() or not f.cleaned_data or f.cleaned_data.get('DELETE'):
                continue
            s = f.instance
            if s.pk:
                continue  # el historial es de solo lectura
            # Nuevo contacto: pasa por el gestor (avanza el embudo y cierra tareas automáticas).
            nuevo = seguimiento.registrar(form.instance, s.tipo, s.notas, request.user, s.fecha)
            if s.vence_en:
                nuevo.vence_en = s.vence_en
                nuevo.save()

    # --- acciones
    def get_actions(self, request, action_location=ActionLocation.CHANGE_LIST):
        acciones = super().get_actions(request, action_location)
        if not request.user.has_perm('prospectos.exportar_prospectos'):
            acciones.pop('exportar_excel', None)
            acciones.pop('exportar_excel_contacto', None)
        elif not request.user.has_perm('prospectos.ver_datos_personales'):
            acciones.pop('exportar_excel_contacto', None)
        if not request.user.has_perm('prospectos.change_prospecto'):
            acciones.pop('asignarme', None)
            acciones.pop('recalcular', None)
        if not (request.user.has_perm('prospectos.change_prospecto') and request.user.has_perm('prospectos.ver_datos_personales')):
            acciones.pop('anonimizar_arco', None)
        return acciones

    @admin.action(description='Anonimizar (derecho de cancelación ARCO, IRREVERSIBLE)')
    def anonimizar_arco(self, request, queryset):
        from django.template.response import TemplateResponse

        from .servicios import seguridad
        if request.POST.get('confirmar') == 'ANONIMIZAR':
            n = 0
            for p in queryset.filter(anonimizado_en__isnull=True):
                seguridad.anonimizar(p, 'derecho_arco')
                n += 1
            auditoria.registrar('anonimizar_arco', None, {'prospectos': n}, request=request)
            self.message_user(request, f'{n} prospectos anonimizados.')
            return None
        return TemplateResponse(request, 'admin/prospectos/prospecto/confirmar_anonimizar.html',
                                {**self.admin_site.each_context(request), 'title': 'Confirmar anonimización',
                                 'prospectos': queryset, 'opts': self.model._meta})

    @admin.action(description='Exportar a Excel (sin datos de contacto)')
    def exportar_excel(self, request, queryset):
        return exportar(request, queryset, con_contacto=False)

    @admin.action(description='Exportar a Excel CON datos de contacto (queda en bitácora)')
    def exportar_excel_contacto(self, request, queryset):
        return exportar(request, queryset, con_contacto=True)

    @admin.action(description='Asignarme como asesor')
    def asignarme(self, request, queryset):
        n = queryset.update(asesor=request.user)
        auditoria.registrar('asignar', None, {'prospectos': n}, request=request)
        self.message_user(request, f'{n} prospectos asignados a ti.')

    @admin.action(description='Recalcular score y segmentos')
    def recalcular(self, request, queryset):
        for p in queryset:
            inteligencia.recalcular(p)
        self.message_user(request, f'{queryset.count()} prospectos recalculados.')


def exportar(request, queryset, con_contacto: bool):
    from io import BytesIO

    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Prospectos'
    columnas = ['ID', 'Aspirante', 'Preparatoria', 'Evento', 'Carrera (alumno)', 'Carrera (tutor)', 'Prioridad',
                'Grupos', 'Estatus', 'Diferencias alumno/tutor', 'Alumno+tutor', 'Registro', 'Último contacto', 'Asesor']
    if con_contacto:
        columnas += ['Correo', 'Teléfono', 'Tutor', 'Correo tutor', 'Teléfono tutor']
    ws.append(columnas)
    qs = queryset.select_related('preparatoria', 'evento', 'carrera', 'carrera_tutor', 'asesor')
    n = 0
    for p in qs.iterator(chunk_size=500):
        fila = [p.pk, p.nombre_aspirante, str(p.preparatoria or ''), p.evento.codigo if p.evento else '', str(p.carrera or ''),
                str(p.carrera_tutor or ''), p.score, ', '.join(etiqueta_segmento(s) for s in p.segmentos or []),
                p.get_estatus_display(), 'Sí' if p.tiene_brecha else 'No', 'Sí' if p.tiene_alumno and p.tiene_tutor else 'No',
                timezone.localtime(p.creado_en).replace(tzinfo=None),
                timezone.localtime(p.ultimo_contacto_en).replace(tzinfo=None) if p.ultimo_contacto_en else None,
                str(p.asesor or '')]
        if con_contacto:
            fila += [p.correo, p.telefono, p.tutor_nombre, p.tutor_correo, p.tutor_telefono]
        # Evita inyección de fórmulas en Excel (=, +, -, @).
        ws.append([("'" + v) if isinstance(v, str) and v[:1] in '=+-@' else v for v in fila])
        n += 1
    for i in range(1, len(columnas) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = 'A2'
    buf = BytesIO()
    wb.save(buf)
    auditoria.registrar('exportar_con_contacto' if con_contacto else 'exportar', None,
                        {'filas': n, 'con_contacto': con_contacto}, request=request)
    r = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r['Content-Disposition'] = f'attachment; filename="prospectos-{timezone.localdate():%Y%m%d}.xlsx"'
    return r


# --------------------------------------------------------------------------- coincidencias
@admin.register(PosibleCoincidencia)
class PosibleCoincidenciaAdmin(ModelAdmin):
    list_display = ['respuesta_nueva', 'prospecto_existente', 'similitud_pct', 'motivo', 'estado', 'creado_en']
    list_filter = ['estado']
    actions = ['fusionar', 'rechazar']
    readonly_fields = ['prospecto', 'candidato', 'similitud', 'motivo', 'estado', 'revisado_por', 'revisado_en', 'creado_en']

    def get_queryset(self, request):
        return super().get_queryset(request).select_related('prospecto__preparatoria', 'candidato__preparatoria')

    def changelist_view(self, request, extra_context=None):
        if 'estado__exact' not in request.GET and not request.GET.get('_changelist_filters'):
            q = request.GET.copy()
            q['estado__exact'] = 'pendiente'
            return HttpResponseRedirect(f'{request.path}?{q.urlencode()}')
        return super().changelist_view(request, extra_context)

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def _desc(self, p):
        prepa = p.preparatoria.nombre if p.preparatoria else 'sin prepa'
        partes = [x for x in ['alumno' if p.tiene_alumno else '', 'tutor' if p.tiene_tutor else ''] if x]
        return format_html('<a class="font-semibold text-link" href="{}">{}</a><br><small class="text-subtle">{} · {}</small>',
                           reverse('admin:prospectos_prospecto_change', args=[p.pk]), p.nombre_aspirante, prepa,
                           ' + '.join(partes) or '—')

    @display(description='Respuesta nueva')
    def respuesta_nueva(self, obj):
        return self._desc(obj.prospecto)

    @display(description='Prospecto existente')
    def prospecto_existente(self, obj):
        return self._desc(obj.candidato)

    @display(description='Similitud', ordering='similitud')
    def similitud_pct(self, obj):
        pct = round(obj.similitud * 100)
        return format_html('<span class="udc-meta"><b>{} %</b><span class="udc-barra"><i style="width:{}%"></i></span></span>', pct, pct)

    def get_actions(self, request, action_location=ActionLocation.CHANGE_LIST):
        a = super().get_actions(request, action_location)
        if not request.user.has_perm('prospectos.resolver_coincidencias'):
            a.pop('fusionar', None)
            a.pop('rechazar', None)
        return a

    @admin.action(description='Es la misma persona: fusionar')
    def fusionar(self, request, queryset):
        hechos = 0
        for c in queryset.filter(estado='pendiente'):
            try:
                inteligencia.fusionar(c.candidato, c.prospecto, request.user)
                hechos += 1
            except inteligencia.NoFusionable as e:
                self.message_user(request, f'{c.prospecto} ↔ {c.candidato}: {e}', messages.ERROR)
        self.message_user(request, f'{hechos} fusionadas.')

    @admin.action(description='Son personas distintas: rechazar')
    def rechazar(self, request, queryset):
        for c in queryset.filter(estado='pendiente'):
            inteligencia.rechazar(c, request.user)
        self.message_user(request, 'Coincidencias rechazadas.')


# --------------------------------------------------------------------------- seguridad
@admin.register(Auditoria)
class AuditoriaAdmin(ModelAdmin):
    list_display = ['creado_en', 'usuario', 'accion', 'objeto_tipo', 'objeto_id', 'detalle', 'ip']
    list_filter = ['accion', 'objeto_tipo', 'usuario']
    date_hierarchy = 'creado_en'

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    def has_delete_permission(self, request, obj=None):
        return False


@admin.register(Respaldo)
class RespaldoAdmin(ModelAdmin):
    list_display = ['creado_en', 'archivo', 'tamano', 'estado_label', 'verificado_en', 'mensaje']
    actions = ['verificar']

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    @display(description='Estado', label={'ok': 'success', 'error': 'danger'})
    def estado_label(self, obj):
        return obj.estado

    @display(description='Tamaño')
    def tamano(self, obj):
        return f'{obj.bytes / 1_048_576:.1f} MB' if obj.bytes else '—'

    def changelist_view(self, request, extra_context=None):
        if request.method == 'POST' and request.POST.get('crear_respaldo') and request.user.is_superuser:
            r = respaldos.crear()
            auditoria.registrar('respaldo_crear', r, request=request)
            self.message_user(request, f'Respaldo {r.archivo}: {r.estado}', messages.SUCCESS if r.estado == 'ok' else messages.ERROR)
            return HttpResponseRedirect(request.path)
        extra = (extra_context or {}) | {'carpeta_respaldos': respaldos.directorio()}
        return super().changelist_view(request, extra)

    @admin.action(description='Verificar (prueba de restauración)')
    def verificar(self, request, queryset):
        for r in queryset.filter(estado='ok'):
            try:
                conteos = respaldos.verificar(r.archivo)
                self.message_user(request, f'{r.archivo}: OK · ' + ', '.join(f'{k.split("_", 1)[-1]}={v}' for k, v in conteos.items()))
            except Exception as e:  # noqa: BLE001
                self.message_user(request, f'{r.archivo}: {e}', messages.ERROR)
